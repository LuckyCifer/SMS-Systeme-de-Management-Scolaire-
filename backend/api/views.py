"""
views.py — SMS (School Management System)
ViewSets complets pour toutes les tables + JWT + Audit + Export CSV
"""
from rest_framework import viewsets, status, filters, serializers
from drf_spectacular.utils import extend_schema, inline_serializer
from django.db.models import Avg, Min, Max, Subquery, OuterRef, Count, Sum, Q, F
from django.db.models.functions import Coalesce
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated, AllowAny, BasePermission
from rest_framework.views import APIView
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from django_filters.rest_framework import DjangoFilterBackend
from django.contrib.auth.models import User
from django.contrib.auth.hashers import check_password
from django.db import IntegrityError
from django.db.models import RestrictedError, ProtectedError
from rest_framework_simplejwt.tokens import RefreshToken
from rest_framework_simplejwt.exceptions import TokenError
from django.http import HttpResponse
from django.core.mail import send_mail
from django.conf import settings as django_settings
import csv
import io
import logging
from datetime import datetime

logger = logging.getLogger(__name__)

from .utils import get_client_ip, log_action
from . import vie_scolaire
from . import horaires
from .notes import (calculer_resultats, deliberer, generer_sequences, series_du_niveau, importer_fiche,
                    periode_close, verifier_periode_ouverte)
from .permissions import RoleBasedPermission
from .recherche import RechercheAvecLibelles
from .mixins import (
    EtablissementFilterMixin, EtablissementParRelationMixin, trusted_etablissement_header, type_etab_courant,
    is_super_admin, get_utilisateur_profile, is_etudiant, get_etudiant_profile, scope_queryset_to_etudiant,
    scope_queryset_to_enseignant, scope_queryset_to_enseignant_cours, assert_enseignant_teaches,
    is_enseignant, get_enseignant_profile,
)
from .models import (
    TypeEtab, Batiment, Salle, Jour, Langue, Module, Pension, Mention,
    TypeEvaluation, Rapport, Annee, Etablissement, Faculte, Departement,
    Specialite, Cycle, Niveau, Classe, MentionClasse,
    Etudiant, Enseignant, Utilisateur, Tuteur, EtudiantTuteur,
    Personnel, POSTES_SIGNATAIRES,
    Tranche, Frais, Inscription, FraisInscription, Paiement, PaiementSalaire, Moratoire,
    Facture, FactureDetail, RapportFinancier,
    Matiere, Qualification, Cours, Periode, Evaluation,
    FicheNotes, FicheNotesDetail,
    Planning, RapportCours, Seance, Absence,
    Examen, Epreuve, Convocation, RapportStatistique, RapportAssiduite,
    Decision, Diplome, CarteEtudiant, Stage, LettreAdmission,
    BadgeAcces, DocumentGenere, AuditLog,
    NiveauScolaire, ConfigBulletin, CoefficientMatiere, ClotureNotes, Appel, PauseHoraire,
)
from .serializers import (
    TypeEtabSerializer, BatimentSerializer, SalleSerializer, JourSerializer,
    LangueSerializer, ModuleSerializer, PensionSerializer, MentionSerializer,
    TypeEvaluationSerializer, RapportSerializer, AnneeSerializer,
    EtablissementSerializer, FaculteSerializer, DepartementSerializer,
    SpecialiteSerializer, CycleSerializer, NiveauSerializer,
    ClasseSerializer, MentionClasseSerializer,
    EtudiantSerializer, EtudiantPhotoSerializer, EnseignantSerializer,
    UtilisateurSerializer, UtilisateurCreateSerializer,
    TuteurSerializer, EtudiantTuteurSerializer,
    PersonnelSerializer,
    TrancheSerializer, FraisSerializer, InscriptionSerializer,
    FraisInscriptionSerializer, PaiementSerializer, PaiementSalaireSerializer, MoratoireSerializer,
    FactureSerializer, FactureDetailSerializer, RapportFinancierSerializer,
    MatiereSerializer, QualificationSerializer, CoursSerializer,
    PeriodeSerializer, EvaluationSerializer,
    FicheNotesSerializer, FicheNotesDetailSerializer,
    PlanningSerializer, RapportCoursSerializer, SeanceSerializer, AbsenceSerializer,
    ExamenSerializer, EpreuveSerializer, ConvocationSerializer, RapportStatistiqueSerializer,
    RapportAssiduiteSerializer,
    DecisionSerializer, DiplomeSerializer, CarteEtudiantSerializer,
    StageSerializer, LettreAdmissionSerializer, BadgeAccesSerializer,
    DocumentGenereSerializer, AuditLogSerializer,
    NiveauScolaireSerializer, ConfigBulletinSerializer, CoefficientMatiereSerializer, ClotureNotesSerializer, AppelSerializer,
    PauseHoraireSerializer,
)




# ── Permissions personnalisées ────────────────────────────────────────────────
class IsAdminSmsUser(BasePermission):
    """Autorise uniquement les utilisateurs avec le rôle ADMIN."""
    message = "Accès réservé aux administrateurs."

    def has_permission(self, request, view):
        if not request.user or not request.user.is_authenticated:
            return False
        try:
            return Utilisateur.objects.get(login=request.user.username).role == 'ADMIN'
        except Utilisateur.DoesNotExist:
            return False


class IsSuperAdmin(BasePermission):
    """Autorise uniquement le SUPER_ADMIN."""
    message = "Accès réservé au super administrateur."

    def has_permission(self, request, view):
        if not request.user or not request.user.is_authenticated:
            return False
        try:
            return Utilisateur.objects.get(login=request.user.username).role == 'SUPER_ADMIN'
        except Utilisateur.DoesNotExist:
            return False


class IsAdminOrSuperAdmin(BasePermission):
    """Autorise ADMIN et SUPER_ADMIN."""
    message = "Accès réservé aux administrateurs."

    def has_permission(self, request, view):
        if not request.user or not request.user.is_authenticated:
            return False
        try:
            role = Utilisateur.objects.get(login=request.user.username).role
            return role in ('ADMIN', 'SUPER_ADMIN')
        except Utilisateur.DoesNotExist:
            return False


# ── Mixin commun ──────────────────────────────────────────────────────────────
def _scoped_queryset(viewset_cls, request):
    """Queryset d'un autre ViewSet, restreint aux droits de l'utilisateur courant."""
    view = viewset_cls(request=request, action='list', format_kwarg=None, kwargs={})
    return view.get_queryset()


class SMSBaseViewSet(viewsets.ModelViewSet):
    permission_classes = [IsAuthenticated, RoleBasedPermission]
    filter_backends    = [DjangoFilterBackend, RechercheAvecLibelles, filters.OrderingFilter]
    # Compte ÉLÈVE : `etudiant_lookup` = chemin vers l'élève (il ne voit que ses lignes) ;
    # `interdit_etudiant` = données réservées au personnel (aucune ligne).
    etudiant_lookup   = None
    interdit_etudiant = False

    def get_queryset(self):
        qs = super().get_queryset()
        if is_etudiant(self.request):
            if self.interdit_etudiant:
                return qs.none()
            if self.etudiant_lookup:
                return scope_queryset_to_etudiant(self.request, qs, self.etudiant_lookup)
        return qs

    def create(self, request, *args, **kwargs):
        try:
            return super().create(request, *args, **kwargs)
        except IntegrityError as exc:
            msg = str(exc)
            if 'Duplicate entry' in msg or '1062' in msg:
                return Response(
                    {'detail': 'Ce matricule existe déjà. Veuillez en choisir un autre.'},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            return Response(
                {'detail': f'Erreur d\'intégrité des données : {msg}'},
                status=status.HTTP_400_BAD_REQUEST,
            )

    def perform_create(self, serializer):
        instance = serializer.save()
        log_action(self.request, 'CREATE',
                   self.__class__.__name__.replace('ViewSet', ''),
                   getattr(instance, instance._meta.pk.name, ''))

    def perform_update(self, serializer):
        instance = serializer.save()
        log_action(self.request, 'UPDATE',
                   self.__class__.__name__.replace('ViewSet', ''),
                   getattr(instance, instance._meta.pk.name, ''))

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        try:
            self.perform_destroy(instance)
        except (RestrictedError, ProtectedError) as exc:
            objs = (
                getattr(exc, 'restricted_objects', None) or
                getattr(exc, 'protected_objects', None) or []
            )
            linked = ', '.join({obj.__class__.__name__ for obj in list(objs)[:10]})
            return Response(
                {'detail': f'Suppression impossible : des enregistrements liés existent ({linked}).'},
                status=status.HTTP_409_CONFLICT,
            )
        return Response(status=status.HTTP_204_NO_CONTENT)

    def verifier_ecriture(self, instance):
        """Crochet : lever une exception pour refuser la suppression de `instance`."""

    def perform_destroy(self, instance):
        self.verifier_ecriture(instance)
        pk = getattr(instance, instance._meta.pk.name, '')
        log_action(self.request, 'DELETE',
                   self.__class__.__name__.replace('ViewSet', ''), pk)
        instance.delete()

    @action(detail=False, methods=['post'], url_path='bulk-delete')
    def bulk_delete(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return Response({'detail': 'Aucun identifiant fourni.'},
                            status=status.HTTP_400_BAD_REQUEST)
        qs    = self.get_queryset().filter(pk__in=ids)
        for instance in qs:
            self.verifier_ecriture(instance)
        count = qs.count()
        log_action(request, 'BULK_DELETE',
                   self.__class__.__name__.replace('ViewSet', ''),
                   detail=f'{count} éléments supprimés')
        try:
            qs.delete()
        except IntegrityError as exc:
            return Response(
                {'detail': f'Suppression impossible : des enregistrements liés existent. ({exc})'},
                status=status.HTTP_409_CONFLICT,
            )
        return Response({'deleted': count}, status=status.HTTP_200_OK)


# ── Login SMS ─────────────────────────────────────────────────────────────────
@extend_schema(
    request=inline_serializer('SmsLoginRequest', {
        'login':    serializers.CharField(),
        'password': serializers.CharField(),
    }),
    responses=inline_serializer('SmsLoginResponse', {
        'login':         serializers.CharField(),
        'nom_user':      serializers.CharField(),
        'role':          serializers.CharField(),
        'type_etab':     serializers.CharField(allow_null=True),
        'lib_type_etab': serializers.CharField(allow_null=True),
        'etablissement': serializers.CharField(allow_null=True),
        'access':        serializers.CharField(),
        'refresh':       serializers.CharField(),
        'message':       serializers.CharField(),
    }),
)
class SmsLoginView(APIView):
    """Authentification par identifiant/mot de passe SMS — retourne les tokens JWT (access + refresh)."""
    permission_classes = [AllowAny]

    def post(self, request):
        login    = request.data.get('login', '').strip()
        password = request.data.get('password', '').strip()

        if not login or not password:
            return Response({'detail': 'Identifiant et mot de passe requis.'},
                            status=status.HTTP_400_BAD_REQUEST)
        try:
            user_sms = Utilisateur.objects.select_related('etablissement').get(login=login)
        except Utilisateur.DoesNotExist:
            AuditLog.objects.create(
                utilisateur=login, action='LOGIN_FAILED', modele='Utilisateur',
                objet_id=login, detail='Utilisateur inexistant',
                ip_address=get_client_ip(request))
            return Response({'detail': 'Identifiant ou mot de passe incorrect.'},
                            status=status.HTTP_401_UNAUTHORIZED)

        passwd_ok = check_password(password, user_sms.passwd or '')

        if not passwd_ok:
            AuditLog.objects.create(
                utilisateur=login, action='LOGIN_FAILED', modele='Utilisateur',
                objet_id=login, detail='Mot de passe incorrect',
                ip_address=get_client_ip(request), etablissement_id=user_sms.etablissement_id)
            return Response({'detail': 'Identifiant ou mot de passe incorrect.'},
                            status=status.HTTP_401_UNAUTHORIZED)

        django_user, _ = User.objects.get_or_create(
            username=login,
            defaults={'is_active': True, 'is_staff': (user_sms.role == 'ADMIN')}
        )
        if django_user.is_staff != (user_sms.role == 'ADMIN'):
            django_user.is_staff = (user_sms.role == 'ADMIN')
            django_user.save()

        refresh = RefreshToken.for_user(django_user)
        AuditLog.objects.create(
            utilisateur=login, action='LOGIN_SUCCESS', modele='Utilisateur',
            objet_id=login, detail=f'Connexion depuis {get_client_ip(request)}',
            ip_address=get_client_ip(request), etablissement_id=user_sms.etablissement_id)

        return Response({
            'login':         user_sms.login,
            'nom_user':      getattr(user_sms, 'nom_user', login),
            'role':          getattr(user_sms, 'role', 'ETUDIANT').upper(),
            'type_etab':     (user_sms.etablissement.type_etab if user_sms.etablissement else None),
            'lib_type_etab': user_sms.type_etab.lib_type if user_sms.type_etab else None,
            'etablissement': user_sms.etablissement_id,
            'access':        str(refresh.access_token),
            'refresh':       str(refresh),
            'message':       'Connexion réussie.',
        }, status=status.HTTP_200_OK)


# ── Logout SMS ───────────────────────────────────────────────────────────────
class SmsLogoutView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        refresh_token = request.data.get('refresh')
        if not refresh_token:
            return Response({'detail': 'Token de rafraîchissement requis.'},
                            status=status.HTTP_400_BAD_REQUEST)
        try:
            token = RefreshToken(refresh_token)
            token.blacklist()
        except TokenError:
            return Response({'detail': 'Token invalide ou déjà révoqué.'},
                            status=status.HTTP_400_BAD_REQUEST)
        log_action(request, 'LOGOUT', 'Utilisateur', request.user.username)
        return Response({'detail': 'Déconnexion réussie.'}, status=status.HTTP_200_OK)


# ── Audit Log ─────────────────────────────────────────────────────────────────
class AuditLogViewSet(viewsets.ReadOnlyModelViewSet):
    queryset           = AuditLog.objects.all().order_by('-date_action')
    serializer_class   = AuditLogSerializer
    # Journal de sécurité — réservé à ADMIN/SUPER_ADMIN même en lecture (RoleBasedPermission
    # ne restreint que l'écriture, insuffisant ici).
    permission_classes = [IsAuthenticated, IsAdminOrSuperAdmin]
    filter_backends    = [RechercheAvecLibelles, filters.OrderingFilter, DjangoFilterBackend]
    search_fields      = ['utilisateur', 'action', 'modele', 'objet_id', 'detail']

    def get_queryset(self):
        # Traces de l'école courante (voir utils.log_action). Le SUPER_ADMIN voit en plus
        # les traces sans école : connexions échouées de comptes inconnus, et ses propres
        # actions antérieures à l'isolation par établissement.
        etab_id = trusted_etablissement_header(self.request)
        qs = super().get_queryset()
        if is_super_admin(self.request):
            return qs.filter(Q(etablissement_id=etab_id) | Q(etablissement__isnull=True))
        return qs.filter(etablissement_id=etab_id) if etab_id else qs.none()
    filterset_fields   = ['action', 'modele', 'utilisateur']

    @action(detail=False, methods=['get'], url_path='stats')
    def stats(self, request):
        qs = self.get_queryset()
        result = qs.aggregate(
            nb_logins=Count('id', filter=Q(action='LOGIN_SUCCESS')),
            nb_creates=Count('id', filter=Q(action='CREATE')),
            nb_updates=Count('id', filter=Q(action='UPDATE')),
            nb_deletes=Count('id', filter=Q(action__in=['DELETE', 'BULK_DELETE'])),
        )
        return Response(result)

    @action(detail=False, methods=['get'], url_path='export-csv')
    def export_csv(self, request):
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="audit_log.csv"'
        response.write('﻿')
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['Date', 'Utilisateur', 'Action', 'Modèle', 'ID Objet', 'Détail', 'IP'])
        for log in self.filter_queryset(self.get_queryset()):
            writer.writerow([
                log.date_action.strftime('%Y-%m-%d %H:%M:%S') if log.date_action else '',
                log.utilisateur, log.action, log.modele,
                log.objet_id, log.detail, log.ip_address,
            ])
        return response


# ── Tables de référence ───────────────────────────────────────────────────────
class TypeEtabViewSet(SMSBaseViewSet):
    queryset         = TypeEtab.objects.all()
    serializer_class = TypeEtabSerializer
    search_fields    = ['lib_type']

class BatimentViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    queryset         = Batiment.objects.all()
    serializer_class = BatimentSerializer

class SalleViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    queryset         = Salle.objects.all()
    serializer_class = SalleSerializer

class JourViewSet(SMSBaseViewSet):
    queryset         = Jour.objects.all()
    serializer_class = JourSerializer

class LangueViewSet(SMSBaseViewSet):
    queryset         = Langue.objects.all()
    serializer_class = LangueSerializer

class ModuleViewSet(SMSBaseViewSet):
    queryset         = Module.objects.all()
    serializer_class = ModuleSerializer
    search_fields    = ['lib_module']

class PensionViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    queryset         = Pension.objects.all()
    serializer_class = PensionSerializer
    filterset_fields = ['type_etab']

    def get_queryset(self):
        qs = super().get_queryset()
        type_str = _resolve_type_etab_str(self.request)
        if type_str:
            from django.db.models import Q
            return qs.filter(Q(type_etab=type_str) | Q(type_etab__isnull=True))
        return qs

class MentionViewSet(SMSBaseViewSet):
    queryset         = Mention.objects.all()
    serializer_class = MentionSerializer

class TypeEvaluationViewSet(SMSBaseViewSet):
    queryset         = TypeEvaluation.objects.all()
    serializer_class = TypeEvaluationSerializer
    filterset_fields = ['type_etab']

    def get_queryset(self):
        qs = super().get_queryset()
        type_str = _resolve_type_etab_str(self.request)
        if type_str:
            from django.db.models import Q
            return qs.filter(Q(type_etab=type_str) | Q(type_etab__isnull=True))
        return qs

class RapportViewSet(SMSBaseViewSet):
    queryset         = Rapport.objects.all()
    serializer_class = RapportSerializer


# ── Année scolaire ────────────────────────────────────────────────────────────
class AnneeViewSet(SMSBaseViewSet):
    queryset         = Annee.objects.all()
    serializer_class = AnneeSerializer
    search_fields    = ['lib_annee', 'code_annee']
    filterset_fields = ['statut']

    @action(detail=False, methods=['get'], url_path='en-cours')
    def annee_en_cours(self, request):
        annee = Annee.objects.filter(statut='EN COURS').first()
        if annee:
            return Response(AnneeSerializer(annee).data)
        return Response({'detail': 'Aucune année en cours.'}, status=status.HTTP_404_NOT_FOUND)

    @action(detail=False, methods=['post'], url_path='passage-annee',
            permission_classes=[IsAuthenticated, IsAdminOrSuperAdmin])
    def passage_annee(self, request):
        """
        Clôture l'année EN COURS et crée la suivante (PLANIFIEE, sans dates).
        Réservé aux rôles ADMIN et SUPER_ADMIN.
        """
        annee_courante = Annee.objects.filter(statut='EN COURS').first()
        if not annee_courante:
            return Response(
                {'detail': 'Aucune année scolaire en cours. Impossible de lancer le passage d\'année.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # ── Calcul du code de la nouvelle année ──────────────────────────────
        code = annee_courante.code_annee.strip()
        try:
            if '-' in code:
                parts = code.split('-')
                # Supporte YYYY-YYYY et YY-YY
                new_parts = [str(int(p) + 1).zfill(len(p)) for p in parts]
                new_code  = '-'.join(new_parts)
            else:
                new_code = str(int(code) + 1)
        except (ValueError, AttributeError):
            return Response(
                {'detail': f"Format du code année '{code}' non reconnu. Attendu : YYYY-YYYY ou YYYY."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if Annee.objects.filter(code_annee=new_code).exists():
            return Response(
                {'detail': f"L'année '{new_code}' existe déjà."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # ── Clôturer l'année courante ─────────────────────────────────────────
        annee_courante.statut = 'CLOTUREE'
        annee_courante.save(update_fields=['statut'])

        # ── Créer la nouvelle année (PLANIFIEE, sans dates) ───────────────────
        new_lib   = f'Année {new_code}'
        annee_new = Annee.objects.create(
            code_annee=new_code,
            lib_annee=new_lib,
            statut='PLANIFIEE',
        )

        log_action(request, 'CREATE', 'Annee', new_code,
                   detail=f"Passage d'année depuis {code} vers {new_code}")

        return Response({
            'annee_cloturee': AnneeSerializer(annee_courante).data,
            'annee_cree':     AnneeSerializer(annee_new).data,
            'message': (
                f"Passage d'année effectué : '{code}' est clôturée. "
                f"'{new_code}' a été créée (statut : PLANIFIÉE). "
                f"Veuillez définir ses dates de début et de fin."
            ),
        }, status=status.HTTP_201_CREATED)


# ── Structure académique ──────────────────────────────────────────────────────
class EtablissementViewSet(SMSBaseViewSet):
    queryset         = Etablissement.objects.select_related('code_type').all()
    serializer_class = EtablissementSerializer
    parser_classes   = [MultiPartParser, FormParser, JSONParser]
    search_fields    = ['lib_etab', 'email']
    filterset_fields = ['code_type', 'type_etab', 'actif', 'statut_agrement']

    def get_permissions(self):
        # Lecture : tout utilisateur authentifié (nécessaire pour afficher son propre
        # établissement dans l'UI). Écriture : ADMIN/SUPER_ADMIN uniquement — sinon
        # n'importe quel rôle (même ETUDIANT) pouvait modifier les infos de l'établissement.
        write_actions = {'create', 'update', 'partial_update', 'destroy'}
        if self.action in write_actions:
            return [IsAuthenticated(), IsAdminOrSuperAdmin()]
        if self.action == 'current' and self.request.method in ('PUT', 'PATCH'):
            return [IsAuthenticated(), IsAdminOrSuperAdmin()]
        return [IsAuthenticated()]

    def get_queryset(self):
        # Seul le SUPER_ADMIN navigue entre les écoles ; tout autre compte est lié à la sienne.
        qs = super().get_queryset()
        if is_super_admin(self.request):
            return qs
        profile = get_utilisateur_profile(self.request)
        return qs.filter(pk=profile.etablissement_id) if profile and profile.etablissement_id else qs.none()

    def partial_update(self, request, *args, **kwargs):
        if request.data.get('clear_logo') in ('1', 'true'):
            instance = self.get_object()
            if instance.logo:
                instance.logo.delete(save=False)
            instance.logo = None
            instance.save(update_fields=['logo'])
            return Response(self.get_serializer(instance).data)
        return super().partial_update(request, *args, **kwargs)

    @action(detail=False, methods=['get', 'put', 'patch'], url_path='current')
    def current(self, request):
        # École courante (mixins.trusted_etablissement_header) : la même que celle qui filtre
        # toutes les données de la requête — jamais une autre, pour que libellés et données
        # correspondent toujours.
        etab = None
        etab_id = trusted_etablissement_header(request)
        if etab_id:
            etab = Etablissement.objects.filter(code_etab=etab_id).first()
        # 4. Dernier recours : n'importe quel établissement (ou création)
        if not etab:
            etab = Etablissement.objects.first()
        if not etab:
            etab = Etablissement.objects.create(
                code_etab='ETAB001',
                lib_etab='Mon Établissement',
                type_etab='SUPERIEUR',
                statut='PRIVE_LAIQUE',
                systeme='FRANCOPHONE',
                actif=True,
            )
        if request.method == 'GET':
            return Response(self.get_serializer(etab).data)
        serializer = self.get_serializer(etab, data=request.data, partial=(request.method == 'PATCH'))
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(serializer.data)

# NEW: Faculté (niveau intermédiaire entre Etablissement et Département)
class FaculteViewSet(EtablissementParRelationMixin, SMSBaseViewSet):
    etablissement_lookups = ('code_etab',)
    queryset         = Faculte.objects.select_related('code_etab').all()
    serializer_class = FaculteSerializer
    search_fields    = ['lib_faculte', 'code_faculte']
    filterset_fields = ['code_etab']

class DepartementViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    queryset         = Departement.objects.select_related('code_etab', 'code_faculte').all()
    serializer_class = DepartementSerializer
    search_fields    = ['lib_dep']
    filterset_fields = ['code_etab', 'code_faculte']

class SpecialiteViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    queryset         = Specialite.objects.select_related('code_dep').all()
    serializer_class = SpecialiteSerializer
    search_fields    = ['lib_sp']
    filterset_fields = ['code_dep']

class CycleViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    queryset         = Cycle.objects.select_related('code_pension').all()
    serializer_class = CycleSerializer
    filterset_fields = ['type_etab']

    def get_queryset(self):
        qs = super().get_queryset()
        type_str = _resolve_type_etab_str(self.request)
        if type_str:
            from django.db.models import Q
            return qs.filter(Q(type_etab=type_str) | Q(type_etab__isnull=True))
        return qs

class NiveauViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    queryset         = Niveau.objects.select_related('code_cycle', 'code_pension').all()
    serializer_class = NiveauSerializer
    filterset_fields = ['code_cycle', 'type_etab']

    def get_queryset(self):
        qs = super().get_queryset()
        type_str = _resolve_type_etab_str(self.request)
        if type_str:
            from django.db.models import Q
            return qs.filter(Q(type_etab=type_str) | Q(type_etab__isnull=True))
        return qs

class ClasseViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    queryset         = Classe.objects.select_related(
        'code_dep', 'code_niveau', 'code_bat', 'code_salle', 'code_sp').all()
    serializer_class = ClasseSerializer
    search_fields    = ['lib_classe', 'code_classe']
    filterset_fields = ['code_dep', 'code_sp', 'code_niveau']

    @action(detail=True, methods=['get'], url_path='export-csv')
    def export_csv(self, request, pk=None):
        pk = self.get_object().pk  # contrôle d'accès (école)
        annee = Annee.objects.filter(statut='EN COURS').first()
        inscriptions = Inscription.objects.filter(
            code_classe=pk, code_annee=annee
        ).select_related('mle_etudiant') if annee else []
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="liste_{pk}.csv"'
        response.write('﻿')
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['Matricule', 'Nom', 'Prénom', 'Sexe', 'Date naissance', 'Nationalité', 'Téléphone', 'Email'])
        for insc in inscriptions:
            e = insc.mle_etudiant
            writer.writerow([
                e.mle_etudiant, e.nom, e.prenom or '',
                e.get_sexe_display() if e.sexe else '',
                e.date_naiss.strftime('%d/%m/%Y') if e.date_naiss else '',
                e.nationalite or '',
                f'\t{e.tel}' if e.tel else '', e.email or '',
            ])
        return response

class MentionClasseViewSet(EtablissementParRelationMixin, SMSBaseViewSet):
    etablissement_lookups = ('code_classe__etablissement',)
    queryset         = MentionClasse.objects.select_related('code_classe', 'code_mention').all()
    serializer_class = MentionClasseSerializer
    filterset_fields = ['code_classe', 'code_mention']


# ── Personnes ─────────────────────────────────────────────────────────────────
class EtudiantViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    etudiant_lookup = 'mle_etudiant'
    queryset         = Etudiant.objects.select_related('code_dep', 'code_sp').all()
    serializer_class = EtudiantSerializer
    search_fields    = ['nom', 'prenom', 'mle_etudiant', 'email', 'tel', 'numero_cni', 'lieu',
                        'nationalite', 'region_or', 'code_dep__lib_dep', 'code_sp__lib_sp']
    filterset_fields = ['code_dep', 'code_sp', 'sexe', 'nationalite']
    parser_classes   = [MultiPartParser, FormParser, JSONParser]

    def get_queryset(self):
        qs = super().get_queryset()
        niveau_sq = Inscription.objects.filter(
            mle_etudiant_id=OuterRef('mle_etudiant')
        ).order_by('-code_annee_id').values(
            'code_classe__code_niveau__lib_niveau'
        )[:1]
        qs = qs.annotate(lib_niveau=Subquery(niveau_sq))
        niveau = self.request.query_params.get('code_niveau')
        if niveau:
            qs = qs.filter(
                inscription__code_classe__code_niveau=niveau
            ).distinct()
        code_classe = self.request.query_params.get('code_classe')
        if code_classe:
            qs = qs.filter(
                inscription__code_classe=code_classe
            ).distinct()
        return qs

    @action(detail=False, methods=['get'], url_path='export-csv')
    def export_csv(self, request):
        qs = self.filter_queryset(self.get_queryset())
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="etudiants.csv"'
        response.write('﻿')
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['Matricule', 'Nom', 'Prénom', 'Sexe', 'Date naissance',
                         'Lieu', 'Région', 'Nationalité', 'N° CNI',
                         'Téléphone', 'Email', 'Filière', 'Spécialité'])
        for e in qs:
            writer.writerow([
                e.mle_etudiant, e.nom, e.prenom or '',
                e.get_sexe_display() if e.sexe else '',
                e.date_naiss.strftime('%d/%m/%Y') if e.date_naiss else '',
                e.lieu or '', e.region_or or '',
                e.nationalite or '', e.numero_cni or '',
                f'\t{e.tel}' if e.tel else '', e.email or '',
                e.code_dep.lib_dep if e.code_dep else '',
                e.code_sp.lib_sp  if e.code_sp  else '',
            ])
        return response

    @action(detail=False, methods=['get'], url_path='export-xlsx')
    def export_xlsx(self, request):
        """Export Excel mis en forme — liste des élèves (inspiré du bulletin)."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        import io as _io

        qs = self.filter_queryset(self.get_queryset())

        # ── Établissement ────────────────────────────────────────────────────
        etab_name, type_etab = '', ''
        etab_id = trusted_etablissement_header(request)
        if etab_id:
            try:
                from .models import Etablissement as _Etab
                etab = _Etab.objects.get(pk=etab_id)
                etab_name  = etab.lib_etab or ''
                type_etab  = etab.type_etab or ''
            except Exception:
                pass

        # ── Couleurs ACERFI ──────────────────────────────────────────────────
        BF, BM, BC = '1F3864', '2E74B5', 'D6E4F0'

        def side(style='thin', color='BBBBBB'):
            return Side(style=style, color=color)
        def thin_b():
            s = side()
            return Border(left=s, right=s, top=s, bottom=s)
        def thick_b():
            s = side('medium', BM)
            return Border(left=s, right=s, top=s, bottom=s)

        HEADERS = ['N°', 'Matricule', 'Nom', 'Prénom', 'Sexe',
                   'Date naissance', 'Lieu', 'Nationalité', 'N° CNI',
                   'Filière / Département', 'Spécialité', 'Niveau',
                   'Téléphone', 'Email']
        COL_W   = [5, 18, 20, 18, 7, 14, 14, 16, 18, 22, 22, 16, 18, 26]
        n       = len(HEADERS)
        lc      = get_column_letter(n)

        wb = Workbook()
        ws = wb.active
        ws.title = 'Liste des élèves'

        # ── En-tête établissement (lignes 1-3) ───────────────────────────────
        ws.row_dimensions[1].height = 14
        ws.merge_cells(f'A1:{lc}1')
        c = ws['A1']
        c.value     = 'REPUBLIQUE DU CAMEROUN  —  Paix - Travail - Patrie'
        c.font      = Font(name='Calibri', size=9, italic=True, color='FFFFFF')
        c.fill      = PatternFill('solid', fgColor=BF)
        c.alignment = Alignment(horizontal='center', vertical='center')

        ws.row_dimensions[2].height = 36
        ws.merge_cells(f'A2:{lc}2')
        c = ws['A2']
        c.value     = etab_name.upper() or 'ÉTABLISSEMENT'
        c.font      = Font(name='Calibri', bold=True, size=16, color='FFFFFF')
        c.fill      = PatternFill('solid', fgColor=BF)
        c.alignment = Alignment(horizontal='center', vertical='center')

        ws.row_dimensions[3].height = 18
        ws.merge_cells(f'A3:{lc}3')
        c = ws['A3']
        c.value     = type_etab
        c.font      = Font(name='Calibri', size=10, color='FFFFFF')
        c.fill      = PatternFill('solid', fgColor=BM)
        c.alignment = Alignment(horizontal='center', vertical='center')

        # ── Titre liste (lignes 4-5) ─────────────────────────────────────────
        ws.row_dimensions[4].height = 10  # séparateur
        ws.row_dimensions[5].height = 30
        ws.merge_cells(f'A5:{lc}5')
        c = ws['A5']
        c.value     = 'LISTE DES ÉLÈVES / ÉTUDIANTS'
        c.font      = Font(name='Calibri', bold=True, size=14, color=BF)
        c.alignment = Alignment(horizontal='center', vertical='center')

        ws.row_dimensions[6].height = 16
        ws.merge_cells(f'A6:{lc}6')
        c = ws['A6']
        c.value     = (f'Exporté le {datetime.now().strftime("%d/%m/%Y à %H:%M")}  —  '
                       f'{qs.count()} enregistrement(s)')
        c.font      = Font(name='Calibri', size=9, italic=True, color='555555')
        c.alignment = Alignment(horizontal='center', vertical='center')

        ws.row_dimensions[7].height = 8  # séparateur

        # ── En-têtes colonnes (ligne 8) ──────────────────────────────────────
        ws.row_dimensions[8].height = 28
        for i, h in enumerate(HEADERS, start=1):
            c = ws.cell(row=8, column=i)
            c.value     = h
            c.font      = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
            c.fill      = PatternFill('solid', fgColor=BF)
            c.alignment = Alignment(horizontal='center', vertical='center',
                                    wrap_text=True)
            s = Side(style='thin', color='FFFFFF')
            c.border    = Border(left=s, right=s, top=s, bottom=s)

        # ── Données ──────────────────────────────────────────────────────────
        alt = [BC, 'FFFFFF']
        for idx, e in enumerate(qs):
            r = 9 + idx
            ws.row_dimensions[r].height = 18
            fill = PatternFill('solid', fgColor=alt[idx % 2])
            row_data = [
                idx + 1,
                e.mle_etudiant,
                e.nom,
                e.prenom or '',
                e.get_sexe_display() if e.sexe else '',
                e.date_naiss.strftime('%d/%m/%Y') if e.date_naiss else '',
                e.lieu or '',
                e.nationalite or '',
                e.numero_cni or '',
                e.code_dep.lib_dep if e.code_dep else '',
                e.code_sp.lib_sp  if e.code_sp  else '',
                getattr(e, 'lib_niveau', '') or '',
                e.tel or '',
                e.email or '',
            ]
            for ci, val in enumerate(row_data, start=1):
                c = ws.cell(row=r, column=ci, value=val)
                c.font      = Font(name='Calibri', size=10)
                c.fill      = fill
                c.alignment = Alignment(vertical='center',
                                        horizontal='center' if ci == 1 else 'left')
                c.border    = thin_b()

        # ── Ligne total ──────────────────────────────────────────────────────
        last_r = 9 + qs.count()
        ws.row_dimensions[last_r].height = 22
        ws.merge_cells(f'A{last_r}:{lc}{last_r}')
        c = ws.cell(row=last_r, column=1)
        c.value     = f'Total : {qs.count()} élève(s) / étudiant(s)'
        c.font      = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
        c.fill      = PatternFill('solid', fgColor=BM)
        c.alignment = Alignment(horizontal='right', vertical='center')

        # ── Auto-filter + freeze ─────────────────────────────────────────────
        ws.auto_filter.ref = f'A8:{lc}8'
        ws.freeze_panes    = 'B9'

        # ── Largeurs colonnes ────────────────────────────────────────────────
        for i, w in enumerate(COL_W, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

        buf = _io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        response = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = 'attachment; filename="liste_etudiants.xlsx"'
        return response

    @action(detail=False, methods=['get'], url_path='template-csv')
    def template_csv(self, request):
        """Télécharger le modèle CSV vide pour l'import d'étudiants."""
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="modele_import_etudiants.csv"'
        response.write('﻿')
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['mle_etudiant', 'nom', 'prenom', 'sexe', 'date_naiss',
                         'lieu', 'nationalite', 'tel', 'email'])
        writer.writerow(['ETU001', 'DUPONT', 'Jean', 'M', '15/03/2000',
                         'Yaoundé', 'Camerounaise', '+237 6XX XXX XXX', 'jean.dupont@email.com'])
        return response

    @action(detail=False, methods=['post'], url_path='import-csv')
    def import_csv(self, request):
        """Importer des étudiants depuis un fichier CSV (;) ou Excel (.xlsx)."""
        file = request.FILES.get('file')
        if not file:
            return Response({'error': 'Aucun fichier fourni.'}, status=status.HTTP_400_BAD_REQUEST)

        filename = file.name.lower()
        created = 0
        updated = 0
        errors  = []

        etab = getattr(request, 'etablissement', None)

        def parse_date(raw):
            for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y'):
                try:
                    return datetime.strptime(raw.strip(), fmt).date()
                except Exception:
                    pass
            return None

        def process_row(line_num, row):
            nonlocal created, updated
            nom = (row.get('nom') or '').strip()
            if not nom:
                errors.append({'ligne': line_num, 'erreur': 'Champ "nom" obligatoire'})
                return
            mle = (row.get('mle_etudiant') or '').strip()
            if not mle:
                errors.append({'ligne': line_num, 'erreur': 'Champ "mle_etudiant" obligatoire'})
                return
            data = {
                'nom':         nom,
                'prenom':      (row.get('prenom') or '').strip() or None,
                'sexe':        (row.get('sexe') or '').strip() or None,
                'date_naiss':  parse_date(row.get('date_naiss') or ''),
                'lieu':        (row.get('lieu') or '').strip() or None,
                'nationalite': (row.get('nationalite') or '').strip() or 'Camerounaise',
                'tel':         (row.get('tel') or '').strip() or None,
                'email':       (row.get('email') or '').strip() or None,
                'nom_tuteur':  (row.get('nom_tuteur') or '').strip() or '',
            }
            if etab:
                data['etablissement'] = etab
            try:
                obj, is_new = Etudiant.objects.get_or_create(mle_etudiant=mle, defaults=data)
                if is_new:
                    created += 1
                else:
                    for k, v in data.items():
                        if v not in (None, ''):
                            setattr(obj, k, v)
                    obj.save()
                    updated += 1
            except Exception as e:
                errors.append({'ligne': line_num, 'erreur': str(e)})

        try:
            if filename.endswith('.csv'):
                content = file.read().decode('utf-8-sig')
                reader  = csv.DictReader(io.StringIO(content), delimiter=';')
                for i, row in enumerate(reader):
                    process_row(i + 2, row)
            elif filename.endswith('.xlsx'):
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(file, data_only=True)
                    ws = wb.active
                    headers = [str(c.value or '').strip() for c in ws[1]]
                    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
                        d = {headers[j]: (str(v).strip() if v is not None else '')
                             for j, v in enumerate(row) if j < len(headers)}
                        process_row(i + 2, d)
                except ImportError:
                    return Response(
                        {'error': 'openpyxl non installé. Utilisez un fichier CSV ou installez openpyxl.'},
                        status=status.HTTP_400_BAD_REQUEST,
                    )
            else:
                return Response({'error': 'Format non supporté (.csv ou .xlsx requis)'},
                                status=status.HTTP_400_BAD_REQUEST)
        except Exception as e:
            return Response({'error': f'Erreur lecture fichier : {e}'}, status=status.HTTP_400_BAD_REQUEST)

        log_action(request, 'CREATE', 'Etudiant', None,
                   f'Import CSV : {created} créés, {updated} MàJ, {len(errors)} erreur(s)')
        return Response({'created': created, 'updated': updated, 'errors': errors})

    @action(detail=True, methods=['get'], url_path='photo')
    def photo(self, request, pk=None):
        serializer = EtudiantPhotoSerializer(
            self.get_object(), context={'request': request}
        )
        return Response(serializer.data)

    @action(detail=True, methods=['delete'], url_path='photo')
    def delete_photo(self, request, pk=None):
        instance = self.get_object()
        if instance.photo:
            instance.photo.delete(save=False)
            instance.photo = None
            instance.save(update_fields=['photo'])
            log_action(request, 'UPDATE', 'Etudiant', pk, 'Photo supprimée')
            return Response({'message': 'Photo supprimée.'})
        return Response({'message': 'Aucune photo.'})

    @action(detail=True, methods=['get'], url_path='inscriptions')
    def inscriptions(self, request, pk=None):
        pk = self.get_object().pk  # contrôle d'accès (école, compte élève)
        qs = Inscription.objects.filter(mle_etudiant=pk).select_related('code_classe', 'code_annee')
        return Response(InscriptionSerializer(qs, many=True).data)

    @action(detail=True, methods=['get'], url_path='paiements')
    def paiements(self, request, pk=None):
        pk = self.get_object().pk  # contrôle d'accès (école, compte élève)
        qs = Paiement.objects.filter(mle_etudiant=pk).select_related('code_tranche', 'code_annee')
        return Response(PaiementSerializer(qs, many=True).data)

    @action(detail=True, methods=['get'], url_path='evaluations')
    def evaluations(self, request, pk=None):
        pk = self.get_object().pk  # contrôle d'accès (école, compte élève)
        qs = Evaluation.objects.filter(mle_etudiant=pk).select_related(
            'code_matiere', 'code_classe', 'code_periode', 'code_type_eval')
        return Response(EvaluationSerializer(qs, many=True).data)

    @action(detail=True, methods=['get'], url_path='tuteurs')
    def tuteurs(self, request, pk=None):
        pk = self.get_object().pk  # contrôle d'accès (école, compte élève)
        qs = EtudiantTuteur.objects.filter(mle_etudiant=pk).select_related('code_tuteur')
        return Response(EtudiantTuteurSerializer(qs, many=True).data)

    @action(detail=True, methods=['get'], url_path='absences')
    def absences(self, request, pk=None):
        pk = self.get_object().pk  # contrôle d'accès (école, compte élève)
        qs = Absence.objects.filter(mle_etudiant=pk).select_related('code_seance')
        return Response(AbsenceSerializer(qs, many=True).data)

    @action(detail=True, methods=['get'], url_path='absences-stats')
    def absences_stats(self, request, pk=None):
        pk = self.get_object().pk  # contrôle d'accès (école, compte élève)
        """Retourne les heures d'absence totales et par matière pour un étudiant."""
        absences = Absence.objects.filter(
            mle_etudiant=pk, present=False
        ).select_related('code_seance__code_matiere')

        total_heures = 0
        par_matiere  = {}

        for absence in absences:
            seance = absence.code_seance
            if not seance:
                continue
            heures = float(seance.nb_heures_effectuees or 0)
            total_heures += heures
            lib = seance.code_matiere.lib_matiere if seance.code_matiere else 'Inconnue'
            if lib not in par_matiere:
                par_matiere[lib] = 0
            par_matiere[lib] += heures

        return Response({
            'mle_etudiant':   pk,
            'total_heures':   round(total_heures, 2),
            'nb_absences':    absences.count(),
            'par_matiere':    [
                {'matiere': k, 'heures': round(v, 2)}
                for k, v in sorted(par_matiere.items(), key=lambda x: -x[1])
            ],
        })

    @action(detail=True, methods=['get'], url_path='decisions')
    def decisions(self, request, pk=None):
        pk = self.get_object().pk  # contrôle d'accès (école, compte élève)
        qs = Decision.objects.filter(mle_etudiant=pk).select_related('code_annee', 'code_classe')
        return Response(DecisionSerializer(qs, many=True).data)

class EnseignantViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    queryset         = Enseignant.objects.select_related('code_dep').all()
    serializer_class = EnseignantSerializer
    search_fields    = ['nom_ens', 'prenom_ens', 'mle_ens', 'email_ens']
    filterset_fields = ['code_dep', 'statut', 'sexe']
    parser_classes   = [MultiPartParser, FormParser, JSONParser]

    @action(detail=True, methods=['delete'], url_path='photo')
    def delete_photo(self, request, pk=None):
        instance = self.get_object()
        if instance.photo:
            instance.photo.delete(save=False)
            instance.photo = None
            instance.save(update_fields=['photo'])
            log_action(request, 'UPDATE', 'Enseignant', pk, 'Photo supprimée')
            return Response({'message': 'Photo supprimée.'})
        return Response({'message': 'Aucune photo.'})

    @action(detail=False, methods=['get'], url_path='export-csv')
    def export_csv(self, request):
        qs = self.filter_queryset(self.get_queryset())
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="enseignants.csv"'
        response.write('﻿')
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['Matricule', 'Nom', 'Prénom', 'Genre', 'N° CNI', 'Téléphone', 'Email', 'Statut', 'Département'])
        for e in qs:
            writer.writerow([
                e.mle_ens, e.nom_ens, e.prenom_ens or '',
                e.get_sexe_display() if e.sexe else '',
                e.numero_cni or '',
                f'\t{e.tel_ens}' if e.tel_ens else '', e.email_ens or '',
                e.get_statut_display() if e.statut else '',
                e.code_dep.lib_dep if e.code_dep else '',
            ])
        return response

    @action(detail=False, methods=['get'], url_path='export-xlsx')
    def export_xlsx(self, request):
        """Export Excel mis en forme — liste des enseignants."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        import io as _io

        qs = self.filter_queryset(self.get_queryset())

        etab_name = ''
        etab_id = trusted_etablissement_header(request)
        if etab_id:
            try:
                from .models import Etablissement as _Etab
                etab_name = _Etab.objects.get(pk=etab_id).lib_etab or ''
            except Exception:
                pass

        BF, BM, BC = '1F3864', '2E74B5', 'D6E4F0'

        def thin_b():
            s = Side(style='thin', color='BBBBBB')
            return Border(left=s, right=s, top=s, bottom=s)

        HEADERS = ['N°', 'Matricule', 'Nom', 'Prénom', 'Genre',
                   'N° CNI', 'Statut', 'Département', 'Téléphone', 'Email']
        COL_W   = [5, 14, 22, 20, 9, 18, 16, 20, 20, 28]
        n, lc   = len(HEADERS), get_column_letter(len(HEADERS))

        wb = Workbook()
        ws = wb.active
        ws.title = 'Liste des enseignants'

        def header_row(row_num, height, text, font_size, bold, color_bg, color_fg='FFFFFF'):
            ws.row_dimensions[row_num].height = height
            ws.merge_cells(f'A{row_num}:{lc}{row_num}')
            c = ws.cell(row=row_num, column=1)
            c.value = text
            c.font = Font(name='Calibri', bold=bold, size=font_size, color=color_fg)
            c.fill = PatternFill('solid', fgColor=color_bg)
            c.alignment = Alignment(horizontal='center', vertical='center')

        header_row(1, 14, 'REPUBLIQUE DU CAMEROUN  —  Paix - Travail - Patrie',
                   9, False, BF)
        header_row(2, 36, etab_name.upper() or 'ÉTABLISSEMENT', 16, True, BF)
        header_row(5, 28, 'LISTE DES ENSEIGNANTS', 14, True, 'FFFFFF', BF)
        ws.row_dimensions[3].height = 8
        ws.row_dimensions[4].height = 8

        ws.row_dimensions[6].height = 15
        ws.merge_cells(f'A6:{lc}6')
        c = ws['A6']
        c.value = (f'Exporté le {datetime.now().strftime("%d/%m/%Y à %H:%M")}  —  '
                   f'{qs.count()} enseignant(s)')
        c.font = Font(name='Calibri', size=9, italic=True, color='555555')
        c.alignment = Alignment(horizontal='center', vertical='center')

        ws.row_dimensions[7].height = 8
        ws.row_dimensions[8].height = 28
        for i, h in enumerate(HEADERS, start=1):
            c = ws.cell(row=8, column=i)
            c.value = h
            c.font = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor=BF)
            c.alignment = Alignment(horizontal='center', vertical='center')
            s = Side(style='thin', color='FFFFFF')
            c.border = Border(left=s, right=s, top=s, bottom=s)

        alt = [BC, 'FFFFFF']
        for idx, e in enumerate(qs):
            r = 9 + idx
            ws.row_dimensions[r].height = 18
            fill = PatternFill('solid', fgColor=alt[idx % 2])
            row_data = [
                idx + 1, e.mle_ens, e.nom_ens, e.prenom_ens or '',
                e.get_sexe_display() if e.sexe else '',
                e.numero_cni or '',
                e.get_statut_display() if e.statut else '',
                e.code_dep.lib_dep if e.code_dep else '',
                e.tel_ens or '', e.email_ens or '',
            ]
            for ci, val in enumerate(row_data, start=1):
                c = ws.cell(row=r, column=ci, value=val)
                c.font = Font(name='Calibri', size=10)
                c.fill = fill
                c.alignment = Alignment(vertical='center',
                                        horizontal='center' if ci == 1 else 'left')
                c.border = thin_b()

        last_r = 9 + qs.count()
        ws.row_dimensions[last_r].height = 22
        ws.merge_cells(f'A{last_r}:{lc}{last_r}')
        c = ws.cell(row=last_r, column=1)
        c.value = f'Total : {qs.count()} enseignant(s)'
        c.font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=BM)
        c.alignment = Alignment(horizontal='right', vertical='center')

        ws.auto_filter.ref = f'A8:{lc}8'
        ws.freeze_panes = 'B9'
        for i, w in enumerate(COL_W, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

        buf = _io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        response = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = 'attachment; filename="liste_enseignants.xlsx"'
        return response

    @action(detail=True, methods=['get'], url_path='planning')
    def planning(self, request, pk=None):
        pk = self.get_object().pk  # contrôle d'accès (école)
        qs = Planning.objects.filter(code_cours__mle_ens=pk).select_related(
            'code_cours__code_matiere', 'code_cours__code_classe', 'code_jour', 'code_salle')
        return Response(PlanningSerializer(qs, many=True).data)

    @action(detail=True, methods=['get'], url_path='seances')
    def seances(self, request, pk=None):
        pk = self.get_object().pk  # contrôle d'accès (école)
        qs = Seance.objects.filter(mle_ens=pk).select_related('code_matiere', 'code_classe')
        return Response(SeanceSerializer(qs, many=True).data)

class UtilisateurViewSet(viewsets.ModelViewSet):
    queryset = Utilisateur.objects.select_related('type_etab', 'etablissement').all()

    def get_queryset(self):
        qs = super().get_queryset()
        from django.db.models import Q

        # SUPER_ADMIN : comptes de l'école à laquelle il est connecté (+ les SUPER_ADMIN).
        if is_super_admin(self.request):
            etab_id = trusted_etablissement_header(self.request)
            return qs.filter(Q(etablissement_id=etab_id) | Q(role='SUPER_ADMIN'))

        if is_etudiant(self.request):
            return qs.filter(login=self.request.user.username)

        # Utilisateur normal : uniquement les comptes de son propre établissement, dérivé de
        # son profil en base (jamais d'un header client), + les comptes SUPER_ADMIN.
        profile = get_utilisateur_profile(self.request)
        etab_id = profile.etablissement_id if profile else None
        if not etab_id:
            return qs.none()
        return qs.filter(
            Q(etablissement_id=etab_id) | Q(type_etab__isnull=True, role='SUPER_ADMIN')
        )

    def get_serializer_class(self):
        if self.action in ['create', 'update', 'partial_update']:
            return UtilisateurCreateSerializer
        return UtilisateurSerializer

    def get_permissions(self):
        # Lecture : tout utilisateur authentifié
        # Écriture/suppression : ADMIN ou SUPER_ADMIN uniquement
        if self.action in ['create', 'update', 'partial_update', 'destroy']:
            return [IsAuthenticated(), IsAdminOrSuperAdmin()]
        return [IsAuthenticated()]

    def perform_create(self, serializer):
        instance = serializer.save()
        log_action(self.request, 'CREATE', 'Utilisateur', instance.login)

    def perform_update(self, serializer):
        instance = serializer.save()
        log_action(self.request, 'UPDATE', 'Utilisateur', instance.login)

class TuteurViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    interdit_etudiant = True  # données réservées au personnel
    queryset         = Tuteur.objects.all()
    serializer_class = TuteurSerializer
    search_fields    = ['nom', 'prenom', 'tel', 'email']
    filterset_fields = ['lien_parente']

class EtudiantTuteurViewSet(EtablissementParRelationMixin, SMSBaseViewSet):
    etudiant_lookup = 'mle_etudiant'
    etablissement_lookups = ('mle_etudiant__etablissement',)
    queryset         = EtudiantTuteur.objects.select_related('mle_etudiant', 'code_tuteur').all()
    serializer_class = EtudiantTuteurSerializer
    filterset_fields = ['mle_etudiant', 'code_tuteur', 'est_contact_principal']


# ── Personnel administratif et de soutien ─────────────────────────────────────
class PersonnelViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    interdit_etudiant = True  # données réservées au personnel
    queryset         = Personnel.objects.select_related('etablissement', 'utilisateur').all()
    serializer_class = PersonnelSerializer
    filterset_fields = ['poste', 'categorie', 'type_contrat', 'actif']
    search_fields    = ['nom', 'prenom', 'mle_personnel', 'matricule_fonct']
    ordering_fields  = ['categorie', 'poste', 'nom', 'date_embauche']
    parser_classes   = [MultiPartParser, FormParser, JSONParser]

    @action(detail=True, methods=['delete'], url_path='photo')
    def delete_photo(self, request, pk=None):
        instance = self.get_object()
        if instance.photo:
            instance.photo.delete(save=False)
            instance.photo = None
            instance.save(update_fields=['photo'])
            log_action(request, 'UPDATE', 'Personnel', pk, 'Photo supprimée')
            return Response({'message': 'Photo supprimée.'})
        return Response({'message': 'Aucune photo.'})

    @action(detail=False, methods=['get'], url_path='signataires')
    def signataires(self, request):
        """Retourne les membres pouvant signer des documents officiels, filtrés par établissement."""
        qs = self.get_queryset().filter(poste__in=POSTES_SIGNATAIRES, actif=True)
        return Response(self.get_serializer(qs, many=True).data)

    @action(detail=False, methods=['get'], url_path='export-csv')
    def export_csv(self, request):
        qs = self.get_queryset()
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="personnel.csv"'
        response.write('﻿')
        writer = csv.writer(response)
        writer.writerow(['Matricule', 'Nom', 'Prénom', 'Sexe', 'Poste', 'Catégorie',
                         'Contrat', 'Date embauche', 'Téléphone', 'Email',
                         'Matricule fonct.', 'Actif'])
        for p in qs:
            writer.writerow([
                p.mle_personnel, p.nom, p.prenom, p.sexe or '',
                p.get_poste_display(), p.get_categorie_display(),
                p.type_contrat, p.date_embauche or '', p.tel, p.email,
                p.matricule_fonct, 'Oui' if p.actif else 'Non',
            ])
        log_action(request, 'EXPORT_CSV', 'Personnel', '')
        return response


# ── Scolarité & Paiements ─────────────────────────────────────────────────────
class TrancheViewSet(EtablissementParRelationMixin, SMSBaseViewSet):
    etablissement_lookups = ('code_pension__etablissement',)
    queryset         = Tranche.objects.select_related('code_pension').all()
    serializer_class = TrancheSerializer
    filterset_fields = ['code_pension']

    def get_queryset(self):
        qs = super().get_queryset()
        type_str = _resolve_type_etab_str(self.request)
        if type_str:
            from django.db.models import Q
            # Filtre via la pension parente
            return qs.filter(
                Q(code_pension__type_etab=type_str) | Q(code_pension__type_etab__isnull=True)
            )
        return qs

class FraisViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    queryset         = Frais.objects.select_related('code_annee').all()
    serializer_class = FraisSerializer
    search_fields    = ['lib_frais', 'type_frais']
    filterset_fields = ['type_frais', 'code_annee', 'type_etab']

    def get_queryset(self):
        qs = super().get_queryset()
        type_str = _resolve_type_etab_str(self.request)
        if type_str:
            from django.db.models import Q
            return qs.filter(Q(type_etab=type_str) | Q(type_etab__isnull=True))
        return qs
    filterset_fields = ['type_frais', 'code_annee']

class InscriptionViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    etudiant_lookup = 'mle_etudiant'
    queryset         = Inscription.objects.select_related(
        'code_classe__code_niveau', 'mle_etudiant', 'code_annee').all()
    serializer_class = InscriptionSerializer
    search_fields    = ['mle_etudiant__nom', 'mle_etudiant__prenom']
    filterset_fields = ['code_classe', 'code_annee']

    def get_queryset(self):
        qs = super().get_queryset()
        code_dep = self.request.query_params.get('code_dep')
        code_sp  = self.request.query_params.get('code_sp')
        if code_dep:
            qs = qs.filter(mle_etudiant__code_dep=code_dep)
        if code_sp:
            qs = qs.filter(mle_etudiant__code_sp=code_sp)
        return qs

    @action(detail=False, methods=['get'], url_path='stats')
    def stats(self, request):
        qs = self.filter_queryset(self.get_queryset())
        paye_sq = Paiement.objects.filter(
            mle_etudiant=OuterRef('mle_etudiant'),
            code_annee=OuterRef('code_annee'),
            type_paiement='INSCRIPTION',
        ).values('mle_etudiant').annotate(total=Sum('mt_paiement')).values('total')
        qs = qs.annotate(mt_paye=Coalesce(Subquery(paye_sq), 0))
        result = qs.aggregate(
            nb_payes=Count('code_inscription', filter=Q(mt_paye__gt=0, mt_paye__gte=F('mt_inscription'))),
            nb_partiels=Count('code_inscription', filter=Q(mt_paye__gt=0, mt_paye__lt=F('mt_inscription'))),
            nb_attente=Count('code_inscription', filter=Q(mt_paye=0)),
            total_recu=Sum('mt_paye'),
        )
        result['total_recu'] = result['total_recu'] or 0
        return Response(result)

    @action(detail=False, methods=['get'], url_path='export-csv')
    def export_csv(self, request):
        qs = self.filter_queryset(self.get_queryset())
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="inscriptions.csv"'
        response.write('﻿')
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['N°', 'Matricule', 'Nom', 'Prénom', 'Classe', 'Année', 'Date inscription', 'Montant (FCFA)'])
        for i in qs:
            e = i.mle_etudiant
            writer.writerow([
                i.code_inscription,
                e.mle_etudiant,
                e.nom,
                e.prenom or '',
                str(i.code_classe),
                str(i.code_annee),
                i.date_inscription.strftime('%d/%m/%Y') if i.date_inscription else '',
                i.mt_inscription,
            ])
        return response

class FraisInscriptionViewSet(EtablissementParRelationMixin, SMSBaseViewSet):
    etudiant_lookup = 'code_inscription__mle_etudiant'
    etablissement_lookups = ('code_inscription__etablissement',)
    queryset         = FraisInscription.objects.select_related(
        'code_inscription', 'code_frais').all()
    serializer_class = FraisInscriptionSerializer
    filterset_fields = ['code_inscription', 'code_frais']

class PaiementViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    etudiant_lookup = 'mle_etudiant'
    queryset         = Paiement.objects.select_related(
        'mle_etudiant__code_dep', 'mle_etudiant__code_sp',
        'code_tranche', 'code_annee').all()
    serializer_class = PaiementSerializer
    search_fields    = ['code_paiement', 'mle_etudiant__nom', 'mle_etudiant__prenom', 'mle_etudiant__mle_etudiant', 'type_paiement', 'statut',
                        'mode_paiement', 'ref_paiement', 'code_tranche__lib_tranche', 'code_annee__code_annee', 'code_annee__lib_annee']
    filterset_fields = ['code_annee', 'code_tranche', 'statut', 'mle_etudiant', 'type_paiement']

    def get_queryset(self):
        qs = super().get_queryset()
        code_dep = self.request.query_params.get('code_dep')
        code_sp  = self.request.query_params.get('code_sp')
        if code_dep:
            qs = qs.filter(mle_etudiant__code_dep=code_dep)
        if code_sp:
            qs = qs.filter(mle_etudiant__code_sp=code_sp)
        return qs

    @action(detail=False, methods=['get'], url_path='stats')
    def stats(self, request):
        qs = self.filter_queryset(self.get_queryset())
        result = qs.aggregate(
            nb_payes=Count('code_paiement', filter=Q(statut='PAYE')),
            nb_partiels=Count('code_paiement', filter=Q(statut='PARTIEL')),
            nb_impayes=Count('code_paiement', filter=Q(statut='IMPAYE')),
            total_encaisse=Sum('mt_paiement', filter=Q(statut='PAYE')),
            nb_inscriptions=Count('code_paiement', filter=Q(type_paiement='INSCRIPTION')),
        )
        result['total_encaisse'] = result['total_encaisse'] or 0
        return Response(result)

    @action(detail=False, methods=['get'], url_path='export-csv')
    def export_csv(self, request):
        qs = self.filter_queryset(self.get_queryset())
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="paiements.csv"'
        response.write('﻿')
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['N°', 'Étudiant', 'Tranche', 'Montant (FCFA)', 'Date', 'Statut', 'Année'])
        for p in qs:
            writer.writerow([
                p.code_paiement,
                f"{p.mle_etudiant.nom} {p.mle_etudiant.prenom or ''}" if p.mle_etudiant else '',
                p.code_tranche.lib_tranche if p.code_tranche else '',
                p.mt_paiement,
                p.date_paiement.strftime('%Y-%m-%d') if p.date_paiement else '',
                p.get_statut_display() if p.statut else '',
                p.code_annee.lib_annee if p.code_annee else '',
            ])
        return response

class PaiementSalaireViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    interdit_etudiant = True  # données réservées au personnel
    """
    Paiements de salaire du personnel — délibérément séparé de PaiementViewSet (frais
    scolarité/inscription des étudiants), qui suit une logique métier différente.
    """
    queryset         = PaiementSalaire.objects.select_related('enseignant', 'personnel').all()
    serializer_class = PaiementSalaireSerializer
    search_fields    = ['enseignant__nom_ens', 'personnel__nom']
    filterset_fields  = ['mois_paie', 'mode_paiement', 'enseignant', 'personnel']

    @action(detail=False, methods=['get'], url_path='stats')
    def stats(self, request):
        qs = self.filter_queryset(self.get_queryset())
        result = qs.aggregate(
            nb_paiements=Count('code_paiement_salaire'),
            total_verse=Sum('montant'),
        )
        result['total_verse'] = result['total_verse'] or 0
        return Response(result)

    @action(detail=False, methods=['get'], url_path='export-csv')
    def export_csv(self, request):
        qs = self.filter_queryset(self.get_queryset())
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="paiements_salaires.csv"'
        response.write('﻿')
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['N°', 'Bénéficiaire', 'Mois', 'Montant (FCFA)', 'Date', 'Mode', 'Référence'])
        for p in qs:
            beneficiaire = p.enseignant or p.personnel
            writer.writerow([
                p.code_paiement_salaire,
                f"{beneficiaire.nom_ens if p.enseignant else beneficiaire.nom} "
                f"{(beneficiaire.prenom_ens if p.enseignant else beneficiaire.prenom) or ''}" if beneficiaire else '',
                p.mois_paie.strftime('%Y-%m') if p.mois_paie else '',
                p.montant,
                p.date_paiement.strftime('%Y-%m-%d') if p.date_paiement else '',
                p.get_mode_paiement_display(),
                p.ref_paiement,
            ])
        return response

class MoratoireViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    etudiant_lookup = 'mle_etudiant'
    queryset         = Moratoire.objects.select_related('mle_etudiant').all()
    serializer_class = MoratoireSerializer
    filterset_fields = ['mle_etudiant']

class FactureViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    etudiant_lookup = 'mle_etudiant'
    queryset         = Facture.objects.select_related('mle_etudiant').all()
    serializer_class = FactureSerializer
    search_fields    = ['numero_facture', 'mle_etudiant__nom', 'mle_etudiant__prenom', 'mle_etudiant__mle_etudiant', 'statut', 'code_annee__code_annee', 'code_annee__lib_annee']
    filterset_fields = ['statut', 'code_annee', 'mle_etudiant']

    @action(detail=False, methods=['get'], url_path='stats')
    def stats(self, request):
        qs = self.filter_queryset(self.get_queryset())
        result = qs.aggregate(
            total_du=Sum('montant_total'),
            total_paye=Sum('montant_paye'),
            nb_factures=Count('code_facture'),
        )
        result['total_du']      = float(result['total_du']    or 0)
        result['total_paye']    = float(result['total_paye']  or 0)
        result['total_restant'] = result['total_du'] - result['total_paye']
        return Response(result)

class FactureDetailViewSet(EtablissementParRelationMixin, SMSBaseViewSet):
    etudiant_lookup = 'code_facture__mle_etudiant'
    etablissement_lookups = ('code_facture__etablissement',)
    queryset         = FactureDetail.objects.select_related('code_facture').all()
    serializer_class = FactureDetailSerializer
    filterset_fields = ['code_facture', 'type_frais']

class RapportFinancierViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    queryset         = RapportFinancier.objects.all()
    serializer_class = RapportFinancierSerializer
    filterset_fields = ['code_annee', 'genere_par']


# ── Helper : résoudre type_etab string depuis les headers HTTP ────────────────
_VALID_TYPE_ETAB = {'PRIMAIRE', 'SECONDAIRE', 'SUPERIEUR'}

def _resolve_type_etab_str(request):
    """Type (PRIMAIRE/SECONDAIRE/SUPERIEUR) de l'école courante — voir mixins.py."""
    type_etab = type_etab_courant(request)
    return type_etab if type_etab in _VALID_TYPE_ETAB else None


class MatiereViewSet(SMSBaseViewSet):
    queryset         = Matiere.objects.select_related('code_module').all()
    serializer_class = MatiereSerializer
    search_fields    = ['lib_matiere', 'code_matiere']
    filterset_fields = ['code_module', 'type_etab']

    def get_queryset(self):
        qs = super().get_queryset()
        type_str = _resolve_type_etab_str(self.request)
        if type_str:
            from django.db.models import Q
            return qs.filter(Q(type_etab=type_str) | Q(type_etab__isnull=True))
        return qs

class QualificationViewSet(EtablissementParRelationMixin, SMSBaseViewSet):
    etablissement_lookups = ('mle_ens__etablissement',)
    queryset         = Qualification.objects.select_related('code_matiere', 'mle_ens').all()
    serializer_class = QualificationSerializer
    filterset_fields = ['mle_ens', 'code_matiere']

class CoursViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    queryset         = Cours.objects.select_related('code_matiere', 'code_classe', 'mle_ens', 'code_annee').all()
    serializer_class = CoursSerializer
    filterset_fields = ['code_classe', 'code_matiere', 'mle_ens', 'semestre', 'code_annee']

    def get_queryset(self):
        qs = super().get_queryset()
        qs = scope_queryset_to_enseignant(self.request, qs, field='mle_ens')
        code_dep = self.request.query_params.get('code_dep')
        code_sp  = self.request.query_params.get('code_sp')
        if code_dep:
            qs = qs.filter(code_classe__code_dep=code_dep)
        if code_sp:
            qs = qs.filter(code_classe__code_sp=code_sp)
        return qs

class PeriodeViewSet(SMSBaseViewSet):
    queryset         = Periode.objects.select_related('code_annee', 'periode_parent').all()
    serializer_class = PeriodeSerializer
    filterset_fields = ['code_annee', 'type_etab', 'periode_parent']

    @action(detail=False, methods=['post'], url_path='generer-sequences')
    def generer_sequences(self, request):
        """Crée les 6 séquences du secondaire (2 par trimestre) pour l'année donnée."""
        code_annee = request.data.get('code_annee')
        if not code_annee:
            return Response({'error': 'code_annee est requis.'}, status=status.HTTP_400_BAD_REQUEST)
        creees = generer_sequences(code_annee)
        log_action(request, 'CREATE', 'Periode', code_annee, f'{creees} séquences générées')
        return Response({'creees': creees})

    def get_queryset(self):
        qs = super().get_queryset()
        type_str = _resolve_type_etab_str(self.request)
        if type_str:
            from django.db.models import Q
            return qs.filter(Q(type_etab=type_str) | Q(type_etab__isnull=True))
        return qs

class EvaluationViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    etudiant_lookup = 'mle_etudiant'
    queryset         = Evaluation.objects.select_related(
        'mle_etudiant', 'code_matiere', 'code_classe', 'code_periode', 'code_type_eval').all()
    serializer_class = EvaluationSerializer
    search_fields    = ['mle_etudiant__nom', 'mle_etudiant__prenom', 'mle_etudiant__mle_etudiant',
                        'code_matiere__lib_matiere', 'code_classe__lib_classe',
                        'code_type_eval__lib_type_eval', 'appreciation']
    filterset_fields = ['code_classe', 'code_periode', 'code_type_eval', 'mle_etudiant',
                        'code_matiere', 'code_annee']

    def get_queryset(self):
        qs = super().get_queryset()
        qs = scope_queryset_to_enseignant_cours(self.request, qs)
        code_dep = self.request.query_params.get('code_dep')
        code_sp  = self.request.query_params.get('code_sp')
        if code_dep:
            qs = qs.filter(code_classe__code_dep=code_dep)
        if code_sp:
            qs = qs.filter(code_classe__code_sp=code_sp)
        return qs

    def perform_create(self, serializer):
        assert_enseignant_teaches(
            self.request,
            serializer.validated_data['code_matiere'].pk,
            serializer.validated_data['code_classe'].pk,
        )
        verifier_periode_ouverte(serializer.validated_data['code_classe'].pk,
                                 serializer.validated_data['code_periode'].pk)
        super().perform_create(serializer)

    def perform_update(self, serializer):
        ancien       = serializer.instance
        code_matiere = serializer.validated_data.get('code_matiere', ancien.code_matiere)
        code_classe  = serializer.validated_data.get('code_classe',  ancien.code_classe)
        code_periode = serializer.validated_data.get('code_periode', ancien.code_periode)
        assert_enseignant_teaches(self.request, code_matiere.pk, code_classe.pk)
        verifier_periode_ouverte(ancien.code_classe_id, ancien.code_periode_id)
        verifier_periode_ouverte(code_classe.pk, code_periode.pk)
        avant = 'absence justifiée' if ancien.dispense else (ancien.note if ancien.note is not None else ancien.appreciation)
        instance = self._save_dans_etablissement(serializer, 'UPDATE', detail='')
        apres = 'absence justifiée' if instance.dispense else (instance.note if instance.note is not None else instance.appreciation)
        if str(avant) != str(apres):
            log_action(self.request, 'UPDATE', 'Evaluation', instance.pk,
                       f'{instance.mle_etudiant_id} / {instance.code_matiere_id} : {avant} → {apres}')

    def verifier_ecriture(self, instance):
        verifier_periode_ouverte(instance.code_classe_id, instance.code_periode_id)

    @action(detail=False, methods=['get'], url_path='export-csv')
    def export_csv(self, request):
        qs = self.filter_queryset(self.get_queryset())
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="evaluations.csv"'
        response.write('﻿')
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['Matricule', 'Étudiant', 'Matière', 'Classe', 'Type', 'Date', 'Note', 'Appréciation'])
        for e in qs:
            writer.writerow([
                e.mle_etudiant.mle_etudiant if e.mle_etudiant else '',
                str(e.mle_etudiant) if e.mle_etudiant else '',
                e.code_matiere.lib_matiere if e.code_matiere else '',
                e.code_classe.code_classe  if e.code_classe  else '',
                e.code_type_eval.lib_type_eval if e.code_type_eval else '',
                e.date_eval.strftime('%Y-%m-%d') if e.date_eval else '',
                e.note if e.note is not None else '',
                e.get_appreciation_display() if e.appreciation else '',
            ])
        return response

    @action(detail=False, methods=['get'], url_path='moyennes')
    def moyennes(self, request):
        """
        Moyennes, rangs et appréciations d'une classe — voir api/notes.py (source unique
        de ces calculs : coefficients au primaire/secondaire, crédits au supérieur,
        séquences, appréciations A/ECA/NA du primaire, rangs ex-aequo).

        Query params : code_classe (requis) + code_periode OU code_annee (bilan annuel).
        mle_etudiant (optionnel) : ne renvoie que cet élève, son rang restant calculé
        sur toute la classe.
        """
        code_classe  = request.query_params.get('code_classe')
        code_periode = request.query_params.get('code_periode')
        code_annee   = request.query_params.get('code_annee')
        mle_etudiant = request.query_params.get('mle_etudiant')
        if not code_classe or not (code_periode or code_annee):
            return Response(
                {'error': 'code_classe et code_periode (ou code_annee) sont requis.'},
                status=status.HTTP_400_BAD_REQUEST,
            )
        # Les filtres de l'URL (code_periode, mle_etudiant…) ne doivent pas restreindre
        # le calcul : le rang se calcule sur toute la classe et toutes ses séquences.
        base = self.get_queryset()
        if is_etudiant(request):
            etudiant = get_etudiant_profile(request)
            if not etudiant:
                return Response([])
            mle_etudiant = etudiant.pk
            base = Evaluation.objects.filter(etablissement_id=self.get_etablissement_id())
        try:
            resultats = calculer_resultats(base, code_classe, code_periode=code_periode,
                                           code_annee=code_annee)
        except Periode.DoesNotExist:
            return Response({'error': 'Période introuvable.'}, status=status.HTTP_404_NOT_FOUND)
        if mle_etudiant:
            resultats = [r for r in resultats if r['mle_etudiant'] == mle_etudiant]
        return Response(resultats)

ROLES_GESTION_NOTES = {'SUPER_ADMIN', 'ADMIN', 'SCOLARITE', 'CENSEUR'}


def exiger_role_gestion_notes(request, quoi):
    from rest_framework.exceptions import PermissionDenied
    profile = get_utilisateur_profile(request)
    if not profile or profile.role not in ROLES_GESTION_NOTES:
        raise PermissionDenied(f"Seuls le censeur et l'administration peuvent {quoi}.")


class ClotureNotesViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    """
    Clôture des notes d'une classe pour une période (après le conseil de classe) :
    POST = clôturer, DELETE = rouvrir. Réservé au censeur et à l'administration.
    """
    queryset         = ClotureNotes.objects.select_related('code_classe', 'code_periode').all()
    serializer_class = ClotureNotesSerializer
    filterset_fields = ['code_classe', 'code_periode']
    interdit_etudiant = True

    def check_permissions(self, request):
        super().check_permissions(request)
        if request.method not in ('GET', 'HEAD', 'OPTIONS'):
            exiger_role_gestion_notes(request, 'clôturer ou rouvrir une période')

    def perform_create(self, serializer):
        from rest_framework.exceptions import NotFound
        classe = serializer.validated_data['code_classe']
        if not _scoped_queryset(ClasseViewSet, self.request).filter(pk=classe.pk).exists():
            raise NotFound('Classe introuvable.')
        serializer.validated_data['cloturee_par'] = self.request.user.username
        super().perform_create(serializer)
        log_action(self.request, 'UPDATE', 'ClotureNotes', classe.pk,
                   f"Clôture des notes : {classe.pk} / {serializer.validated_data['code_periode']}")

    def perform_destroy(self, instance):
        log_action(self.request, 'UPDATE', 'ClotureNotes', instance.code_classe_id,
                   f'Réouverture des notes : {instance.code_classe_id} / {instance.code_periode}')
        instance.delete()


# ── Vie scolaire : appel journalier, absences du jour, alertes ────────────────
ROLES_VIE_SCOLAIRE = {'SUPER_ADMIN', 'ADMIN', 'SCOLARITE', 'CENSEUR', 'SURVEILLANT_GENERAL', 'DIRECTION'}
ROLES_JUSTIFICATION = {'SUPER_ADMIN', 'ADMIN', 'SCOLARITE', 'CENSEUR', 'SURVEILLANT_GENERAL'}


def _date_param(valeur):
    import datetime
    from rest_framework.exceptions import ValidationError
    if not valeur:
        return datetime.date.today()
    try:
        return datetime.date.fromisoformat(valeur)
    except ValueError:
        raise ValidationError({'error': 'Date invalide (format AAAA-MM-JJ).'})


class AppelViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    """
    Appel journalier d'une classe (matin / après-midi). GET feuille = liste d'appel (présent
    par défaut) ; POST enregistrer = enregistrement tout ou rien. Un enseignant ne fait
    l'appel que dans les classes où il enseigne.
    """
    queryset          = Appel.objects.select_related('code_classe').all()
    serializer_class  = AppelSerializer
    filterset_fields  = ['code_classe', 'date', 'creneau']
    interdit_etudiant = True

    def _classe_autorisee(self, code_classe):
        from rest_framework.exceptions import NotFound, PermissionDenied
        if not _scoped_queryset(ClasseViewSet, self.request).filter(pk=code_classe).exists():
            raise NotFound('Classe introuvable.')
        enseignant = get_enseignant_profile(self.request)
        if is_enseignant(self.request) and not (
                enseignant and Cours.objects.filter(mle_ens=enseignant, code_classe=code_classe).exists()):
            raise PermissionDenied("Vous n'enseignez pas dans cette classe.")

    @action(detail=False, methods=['get'], url_path='feuille')
    def feuille(self, request):
        code_classe = request.query_params.get('code_classe')
        creneau = request.query_params.get('creneau') or 'MATIN'
        if not code_classe:
            return Response({'error': 'code_classe est requis.'}, status=status.HTTP_400_BAD_REQUEST)
        self._classe_autorisee(code_classe)
        return Response(vie_scolaire.feuille_appel(code_classe, _date_param(request.query_params.get('date')), creneau))

    @action(detail=False, methods=['post'], url_path='enregistrer')
    def enregistrer(self, request):
        code_classe = request.data.get('code_classe')
        creneau = request.data.get('creneau') or 'MATIN'
        if not code_classe or creneau not in dict(Appel.CRENEAU_CHOICES):
            return Response({'error': 'code_classe et un créneau valide sont requis.'}, status=status.HTTP_400_BAD_REQUEST)
        self._classe_autorisee(code_classe)
        appel, erreurs = vie_scolaire.enregistrer_appel(
            self.get_etablissement_id(), code_classe, _date_param(request.data.get('date')),
            creneau, request.data.get('lignes') or [], request.user.username)
        if erreurs:
            return Response({'error': ' '.join(erreurs)}, status=status.HTTP_400_BAD_REQUEST)
        nb_abs = appel.lignes.filter(statut='ABSENT').count()
        log_action(request, 'UPDATE', 'Appel', appel.pk, f'Appel {code_classe} {appel.date} {creneau} : {nb_abs} absent(s)')
        return Response({'appel': appel.pk, 'absents': nb_abs})


class VieScolaireViewSet(viewsets.ViewSet):
    """Vue du surveillant général : absences et retards du jour, justification, alertes."""
    permission_classes = [IsAuthenticated]

    def _exiger(self, request, roles):
        from rest_framework.exceptions import PermissionDenied
        profile = get_utilisateur_profile(request)
        if not profile or profile.role not in roles:
            raise PermissionDenied("Votre rôle ne vous autorise pas à accéder à la vie scolaire.")
        etab_id = trusted_etablissement_header(request)
        if not etab_id:
            raise PermissionDenied('Aucun établissement sélectionné.')
        return etab_id

    @action(detail=False, methods=['get'], url_path='jour')
    def jour(self, request):
        etab_id = self._exiger(request, ROLES_VIE_SCOLAIRE)
        date = _date_param(request.query_params.get('date'))
        return Response({'date': date.isoformat(), 'lignes': vie_scolaire.evenements_du_jour(etab_id, date)})

    @action(detail=False, methods=['post'], url_path='justifier')
    def justifier(self, request):
        etab_id = self._exiger(request, ROLES_JUSTIFICATION)
        ok = vie_scolaire.justifier(etab_id, request.data.get('source'), request.data.get('id'),
                                    request.data.get('justifiee', True), request.data.get('motif'))
        if not ok:
            return Response({'error': 'Absence introuvable.'}, status=status.HTTP_404_NOT_FOUND)
        log_action(request, 'UPDATE', 'Absence', request.data.get('id'),
                   f"{'Justifiée' if request.data.get('justifiee', True) else 'Non justifiée'} : {request.data.get('motif') or ''}")
        return Response({'ok': True})

    @action(detail=False, methods=['get'], url_path='alertes')
    def alertes(self, request):
        etab_id = self._exiger(request, ROLES_VIE_SCOLAIRE)
        q = request.query_params
        try:
            seuils = dict(seuil_avertissement=float(q.get('seuil_avertissement', 10)),
                          seuil_conseil=float(q.get('seuil_conseil', 20)),
                          heures_demi_journee=float(q.get('heures_demi_journee', vie_scolaire.HEURES_PAR_DEMI_JOURNEE)))
        except ValueError:
            return Response({'error': 'Seuils invalides.'}, status=status.HTTP_400_BAD_REQUEST)
        return Response(vie_scolaire.alertes(etab_id, q.get('code_annee'), **seuils))


class FicheNotesViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    interdit_etudiant = True  # données réservées au personnel
    queryset         = FicheNotes.objects.select_related(
        'code_matiere', 'code_classe', 'mle_ens', 'code_periode', 'code_type_eval').all()
    serializer_class = FicheNotesSerializer
    search_fields    = ['code_matiere__lib_matiere', 'code_classe__lib_classe', 'mle_ens__nom_ens', 'mle_ens__prenom_ens',
                        'code_type_eval__lib_type_eval', 'code_periode__lib_periode', 'statut']
    filterset_fields = ['code_classe', 'code_annee', 'mle_ens', 'statut']

    def get_queryset(self):
        qs = super().get_queryset()
        # Un compte ENSEIGNANT ne voit/gère que ses propres fiches de notes.
        return scope_queryset_to_enseignant(self.request, qs, field='mle_ens')

    def _assert_mle_ens_is_self(self, mle_ens):
        from rest_framework.exceptions import PermissionDenied
        if not is_enseignant(self.request):
            return
        enseignant = get_enseignant_profile(self.request)
        if not enseignant or mle_ens.pk != enseignant.pk:
            raise PermissionDenied("Vous ne pouvez créer/modifier que vos propres fiches de notes.")

    def perform_create(self, serializer):
        self._assert_mle_ens_is_self(serializer.validated_data['mle_ens'])
        super().perform_create(serializer)

    def perform_update(self, serializer):
        mle_ens = serializer.validated_data.get('mle_ens', serializer.instance.mle_ens)
        self._assert_mle_ens_is_self(mle_ens)
        super().perform_update(serializer)

    # ── Circuit de la fiche ──────────────────────────────────────────────────
    def _verifier_modifiable(self, fiche):
        """Notes modifiables : brouillon (ou fiche soumise, pour le censeur) et période ouverte."""
        from rest_framework.exceptions import PermissionDenied
        profile = get_utilisateur_profile(self.request)
        gestion = bool(profile and profile.role in ROLES_GESTION_NOTES)
        if fiche.statut in ('VALIDE', 'IMPORTE'):
            raise PermissionDenied('Fiche validée : demandez sa réouverture au censeur pour la modifier.')
        if fiche.statut == 'SOUMIS' and not gestion:
            raise PermissionDenied('Fiche soumise au censeur : elle ne peut plus être modifiée par l\'enseignant.')
        verifier_periode_ouverte(fiche.code_classe_id, fiche.code_periode_id)

    @action(detail=True, methods=['post'], url_path='saisir')
    def saisir(self, request, pk=None):
        """
        Enregistre toutes les notes de la fiche en une seule opération (tout ou rien).
        Body : {notes: [{mle_etudiant, note, absent, justifie, observation}]} — note sur le
        barème de la fiche. Les élèves non transmis gardent leur saisie précédente.
        """
        from django.db import transaction
        fiche = self.get_object()
        self._verifier_modifiable(fiche)
        lignes  = request.data.get('notes') or []
        # Élèves inscrits dans la classe POUR L'ANNÉE de la fiche (un élève passé en classe
        # supérieure garde son inscription de l'an dernier dans son ancienne classe).
        inscrits = set(Inscription.objects.filter(code_classe=fiche.code_classe_id, code_annee=fiche.code_annee_id)
                       .values_list('mle_etudiant_id', flat=True))
        erreurs = []
        propres = []
        for l in lignes:
            mle, note = l.get('mle_etudiant'), l.get('note')
            absent = bool(l.get('absent'))
            if mle not in inscrits:
                erreurs.append(f'{mle} : élève non inscrit dans cette classe.')
                continue
            if absent or note in (None, ''):
                note = None
            else:
                try:
                    note = float(note)
                except (TypeError, ValueError):
                    erreurs.append(f'{mle} : note invalide.')
                    continue
                if not 0 <= note <= float(fiche.bareme):
                    erreurs.append(f'{mle} : la note doit être comprise entre 0 et {fiche.bareme}.')
                    continue
            propres.append((mle, note, absent, absent and bool(l.get('justifie')),
                            (l.get('observation') or '')[:100] or None))
        if erreurs:
            return Response({'error': ' '.join(erreurs)}, status=status.HTTP_400_BAD_REQUEST)
        with transaction.atomic():
            for mle, note, absent, justifie, obs in propres:
                FicheNotesDetail.objects.update_or_create(
                    code_fiche=fiche, mle_etudiant_id=mle,
                    defaults={'note': note, 'absent': absent, 'justifie': justifie,
                              'observation': obs, 'appreciation': None},
                )
        log_action(request, 'UPDATE', 'FicheNotes', pk, f'{len(propres)} notes saisies')
        return Response({'enregistrees': len(propres), 'statut': fiche.statut})

    @action(detail=True, methods=['post'], url_path='soumettre')
    def soumettre(self, request, pk=None):
        """L'enseignant transmet sa fiche au censeur ; il ne peut plus la modifier ensuite."""
        fiche = self.get_object()
        if fiche.statut != 'BROUILLON':
            return Response({'error': 'Seule une fiche en brouillon peut être soumise.'},
                            status=status.HTTP_400_BAD_REQUEST)
        fiche.statut = 'SOUMIS'
        fiche.save(update_fields=['statut'])
        log_action(request, 'UPDATE', 'FicheNotes', pk, 'Fiche soumise au censeur')
        return Response({'statut': fiche.statut})

    @action(detail=True, methods=['post'], url_path='valider')
    def valider(self, request, pk=None):
        """Validation par le censeur ou l'administration : les notes sont reportées aussitôt."""
        from django.db import transaction
        exiger_role_gestion_notes(request, 'valider une fiche de notes')
        fiche = self.get_object()
        if fiche.statut in ('VALIDE', 'IMPORTE'):
            return Response({'error': 'Cette fiche est déjà validée.'}, status=status.HTTP_400_BAD_REQUEST)
        verifier_periode_ouverte(fiche.code_classe_id, fiche.code_periode_id)
        with transaction.atomic():
            reportees = importer_fiche(fiche)
            fiche.statut = 'VALIDE'
            fiche.save(update_fields=['statut'])
        log_action(request, 'UPDATE', 'FicheNotes', pk, f'Fiche validée : {reportees} notes reportées')
        return Response({'statut': fiche.statut, 'imported': reportees})

    @action(detail=True, methods=['post'], url_path='rouvrir')
    def rouvrir(self, request, pk=None):
        """Le censeur rouvre une fiche (soumise ou validée) pour correction ; à revalider ensuite."""
        exiger_role_gestion_notes(request, 'rouvrir une fiche de notes')
        fiche = self.get_object()
        verifier_periode_ouverte(fiche.code_classe_id, fiche.code_periode_id)
        fiche.statut = 'BROUILLON'
        fiche.save(update_fields=['statut'])
        log_action(request, 'UPDATE', 'FicheNotes', pk, 'Fiche rouverte pour correction')
        return Response({'statut': fiche.statut})

    @action(detail=True, methods=['post'], url_path='importer')
    def importer_evaluations(self, request, pk=None):
        """Ancien circuit (valider puis importer) : équivalent à la validation."""
        return self.valider(request, pk)

class FicheNotesDetailViewSet(EtablissementParRelationMixin, SMSBaseViewSet):
    def _verifier_fiche(self, fiche):
        # Mêmes règles que la saisie groupée (statut de la fiche, période ouverte).
        FicheNotesViewSet(request=self.request, format_kwarg=None, kwargs={})._verifier_modifiable(fiche)

    def perform_create(self, serializer):
        self._verifier_fiche(serializer.validated_data['code_fiche'])
        super().perform_create(serializer)

    def perform_update(self, serializer):
        self._verifier_fiche(serializer.validated_data.get('code_fiche', serializer.instance.code_fiche))
        super().perform_update(serializer)

    def verifier_ecriture(self, instance):
        self._verifier_fiche(instance.code_fiche)

    etudiant_lookup = 'mle_etudiant'
    etablissement_lookups = ('code_fiche__etablissement',)
    queryset         = FicheNotesDetail.objects.select_related('code_fiche', 'mle_etudiant').all()
    serializer_class = FicheNotesDetailSerializer
    filterset_fields = ['code_fiche', 'mle_etudiant', 'absent']


# ── Planning & Rapports ───────────────────────────────────────────────────────
class PauseHoraireViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    """
    Pauses de l'emploi du temps de l'école connectée. Lecture pour tous ; écriture réservée
    au censeur et à l'administration. Une pause ne peut pas chevaucher un cours programmé.
    """
    queryset         = PauseHoraire.objects.select_related('code_jour').all()
    serializer_class = PauseHoraireSerializer
    pagination_class = None
    filterset_fields = ['code_jour']

    def check_permissions(self, request):
        super().check_permissions(request)
        if request.method not in ('GET', 'HEAD', 'OPTIONS'):
            exiger_role_gestion_notes(request, "modifier les pauses de l'emploi du temps")

    def perform_destroy(self, instance):
        # Supprimer la dernière pause propre à un jour y rétablit les pauses de la semaine :
        # elles ne doivent pas tomber sur un cours de ce jour.
        if instance.code_jour_id and not PauseHoraire.objects.filter(
                etablissement_id=instance.etablissement_id, code_jour_id=instance.code_jour_id,
        ).exclude(pk=instance.pk).exists():
            for p in PauseHoraire.objects.filter(etablissement_id=instance.etablissement_id,
                                                 code_jour__isnull=True):
                cours = horaires.cours_chevauches(
                    instance.etablissement_id, [instance.code_jour_id], p.h_debut, p.h_fin,
                ).filter(type_planning='HEBDO').first()
                if cours:
                    raise serializers.ValidationError({'detail': (
                        f"Sans pause propre au {instance.code_jour.lib_jour}, la pause « {p.libelle} » "
                        f"de la semaine s'y appliquerait. " + horaires.message_cours(cours))})
        super().perform_destroy(instance)


class PlanningViewSet(EtablissementParRelationMixin, SMSBaseViewSet):
    etablissement_lookups = ('code_cours__etablissement',)
    queryset         = Planning.objects.select_related(
        'code_cours__code_matiere', 'code_cours__code_classe',
        'code_cours__mle_ens', 'code_cours__code_annee',
        'code_jour', 'code_salle').all()
    serializer_class = PlanningSerializer
    filterset_fields = ['code_cours__code_classe', 'code_cours__mle_ens',
                        'code_cours__code_annee', 'code_cours__semestre',
                        'code_jour', 'type_planning', 'type_seance']

    @action(detail=False, methods=['get'], url_path='classe/(?P<code_classe>[^/.]+)')
    def par_classe(self, request, code_classe=None):
        qs = Planning.objects.filter(
            code_cours__code_classe=code_classe
        ).select_related(
            'code_cours__code_matiere', 'code_cours__mle_ens',
            'code_cours__code_annee', 'code_jour', 'code_salle'
        ).order_by('code_jour', 'h_debut')
        return Response(PlanningSerializer(qs, many=True).data)

    @action(detail=False, methods=['post'], url_path='generer-seances')
    def generer_seances(self, request):
        """
        Génère des séances de cours pour une semaine donnée à partir
        des créneaux HEBDO du planning.

        Body: { date_lundi: "2025-10-07", code_annee: "2025-2026" }
        """
        from datetime import date, timedelta
        from decimal import Decimal

        date_lundi_str = request.data.get('date_lundi')
        code_annee     = request.data.get('code_annee')

        if not date_lundi_str or not code_annee:
            return Response(
                {'detail': 'date_lundi et code_annee sont requis.'},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            date_lundi = date.fromisoformat(date_lundi_str)
        except ValueError:
            return Response({'detail': 'Format de date invalide (YYYY-MM-DD).'}, status=400)

        # Correspondance jour django → décalage lundi=0
        JOUR_CODE_TO_OFFSET = {
            'LUN': 0, 'MAR': 1, 'MER': 2, 'JEU': 3, 'VEN': 4, 'SAM': 5, 'DIM': 6,
        }

        plannings = self.get_queryset().filter(
            type_planning='HEBDO',
            code_cours__code_annee=code_annee,
        ).select_related(
            'code_cours__code_matiere', 'code_cours__code_classe',
            'code_cours__mle_ens', 'code_jour',
        )

        created = 0
        skipped = 0

        for p in plannings:
            code_jour = p.code_jour.code_jour if p.code_jour else None
            offset    = JOUR_CODE_TO_OFFSET.get(code_jour)
            if offset is None:
                continue

            date_seance = date_lundi + timedelta(days=offset)

            # Éviter les doublons pour le même créneau ce jour-là
            already = Seance.objects.filter(
                code_matiere=p.code_cours.code_matiere,
                code_classe=p.code_cours.code_classe,
                code_annee=code_annee,
                date_seance=date_seance,
                h_debut=p.h_debut,
            ).exists()
            if already:
                skipped += 1
                continue

            h_debut = p.h_debut
            h_fin   = p.h_fin
            nb_h    = Decimal('0')
            if h_debut and h_fin:
                debut_min = h_debut.hour * 60 + h_debut.minute
                fin_min   = h_fin.hour   * 60 + h_fin.minute
                nb_h      = Decimal(str(round((fin_min - debut_min) / 60, 2)))

            Seance.objects.create(
                code_matiere=p.code_cours.code_matiere,
                code_classe=p.code_cours.code_classe,
                mle_ens=p.code_cours.mle_ens,
                code_annee_id=code_annee,
                date_seance=date_seance,
                h_debut=h_debut,
                h_fin=h_fin,
                salle=p.code_salle.code_salle if p.code_salle else '',
                nb_heures_effectuees=nb_h,
                statut='TENU',
            )
            created += 1

        log_action(request, 'CREATE', 'Seance', detail=f'{created} séances générées pour la semaine du {date_lundi_str}')
        return Response({'created': created, 'skipped': skipped}, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['post'], url_path='generer-periode')
    def generer_periode(self, request):
        """
        Génère des séances pour toute une période.
        Body: { date_debut: "YYYY-MM-DD", date_fin: "YYYY-MM-DD", code_annee: "2025-2026" }
        - Créneaux HEBDO    : une séance par semaine pour chaque créneau entre date_debut et date_fin
        - Créneaux INTENSIF : une séance par jour de la plage du slot, intersectée avec la période
        """
        from datetime import date, timedelta
        from decimal import Decimal

        date_debut_str = request.data.get('date_debut')
        date_fin_str   = request.data.get('date_fin')
        code_annee     = request.data.get('code_annee')

        if not date_debut_str or not date_fin_str or not code_annee:
            return Response(
                {'detail': 'date_debut, date_fin et code_annee sont requis.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            date_debut = date.fromisoformat(date_debut_str)
            date_fin   = date.fromisoformat(date_fin_str)
        except ValueError:
            return Response({'detail': 'Format de date invalide (YYYY-MM-DD).'}, status=400)

        if date_fin < date_debut:
            return Response({'detail': 'date_fin doit être postérieure à date_debut.'}, status=400)

        JOUR_CODE_TO_OFFSET = {
            'LUN': 0, 'MAR': 1, 'MER': 2, 'JEU': 3, 'VEN': 4, 'SAM': 5, 'DIM': 6,
        }

        base_qs = self.get_queryset().filter(
            code_cours__code_annee=code_annee,
        ).select_related(
            'code_cours__code_matiere', 'code_cours__code_classe',
            'code_cours__mle_ens', 'code_jour',
        )

        created = 0
        skipped = 0

        def _exists(p, d):
            return Seance.objects.filter(
                code_matiere=p.code_cours.code_matiere,
                code_classe=p.code_cours.code_classe,
                code_annee=code_annee,
                date_seance=d,
                h_debut=p.h_debut,
            ).exists()

        def _create(p, d):
            nonlocal created
            h_debut = p.h_debut
            h_fin   = p.h_fin
            nb_h    = Decimal('0')
            if h_debut and h_fin:
                debut_min = h_debut.hour * 60 + h_debut.minute
                fin_min   = h_fin.hour   * 60 + h_fin.minute
                nb_h      = Decimal(str(round((fin_min - debut_min) / 60, 2)))
            Seance.objects.create(
                code_matiere=p.code_cours.code_matiere,
                code_classe=p.code_cours.code_classe,
                mle_ens=p.code_cours.mle_ens,
                code_annee_id=code_annee,
                date_seance=d,
                h_debut=h_debut,
                h_fin=h_fin,
                salle=p.code_salle.code_salle if p.code_salle else '',
                nb_heures_effectuees=nb_h,
                statut='TENU',
            )
            created += 1

        # ── HEBDO : semaine par semaine ───────────────────────────────────────
        hebdo_qs = list(base_qs.filter(type_planning='HEBDO'))
        if hebdo_qs:
            first_monday  = date_debut - timedelta(days=date_debut.weekday())
            cur_monday    = first_monday
            while cur_monday <= date_fin:
                for p in hebdo_qs:
                    code_jour = p.code_jour.code_jour if p.code_jour else None
                    offset    = JOUR_CODE_TO_OFFSET.get(code_jour)
                    if offset is None:
                        continue
                    d = cur_monday + timedelta(days=offset)
                    if d < date_debut or d > date_fin:
                        continue
                    if _exists(p, d):
                        skipped += 1
                    else:
                        _create(p, d)
                cur_monday += timedelta(weeks=1)

        # ── INTENSIF : une séance par jour dans l'intersection ────────────────
        intensif_qs = base_qs.filter(
            type_planning='INTENSIF',
            date_debut__lte=date_fin,
            date_fin__gte=date_debut,
        )
        for p in intensif_qs:
            cur = max(p.date_debut, date_debut)
            end = min(p.date_fin,   date_fin)
            while cur <= end:
                if _exists(p, cur):
                    skipped += 1
                else:
                    _create(p, cur)
                cur += timedelta(days=1)

        log_action(
            request, 'CREATE', 'Seance',
            detail=f'{created} séances générées du {date_debut_str} au {date_fin_str}',
        )
        return Response({'created': created, 'skipped': skipped}, status=status.HTTP_201_CREATED)

class RapportCoursViewSet(EtablissementParRelationMixin, SMSBaseViewSet):
    etablissement_lookups = ('code_classe__etablissement',)
    queryset         = RapportCours.objects.select_related(
        'code_matiere', 'code_classe', 'code_rapport', 'mle_ens').all()
    serializer_class = RapportCoursSerializer
    filterset_fields = ['code_classe', 'mle_ens', 'code_rapport']

class SeanceViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    queryset         = Seance.objects.select_related(
        'code_matiere', 'code_classe', 'mle_ens', 'code_annee').all()
    serializer_class = SeanceSerializer
    search_fields    = ['code_matiere__lib_matiere', 'code_classe__lib_classe', 'mle_ens__nom_ens', 'mle_ens__prenom_ens',
                        'salle', 'statut']
    filterset_fields = ['code_classe', 'mle_ens', 'code_annee', 'statut', 'date_seance']

    def get_queryset(self):
        qs = super().get_queryset()
        # Un compte ENSEIGNANT ne voit / ne remplit que ses propres séances.
        qs = scope_queryset_to_enseignant(self.request, qs, field='mle_ens')
        code_dep = self.request.query_params.get('code_dep')
        code_sp  = self.request.query_params.get('code_sp')
        if code_dep:
            qs = qs.filter(code_classe__code_dep=code_dep)
        if code_sp:
            qs = qs.filter(code_classe__code_sp=code_sp)
        return qs

    @action(detail=True, methods=['get'], url_path='presences')
    def presences(self, request, pk=None):
        seance   = self.get_object()  # contrôle d'accès (école, enseignant)
        absences = scope_queryset_to_etudiant(
            request, Absence.objects.filter(code_seance=seance).select_related('mle_etudiant'))
        return Response(AbsenceSerializer(absences, many=True).data)

    @action(detail=True, methods=['post'], url_path='saisir-presences')
    def saisir_presences(self, request, pk=None):
        """Crée ou met à jour les présences d'une séance.
        Body: { presences: [{mle_etudiant, present, signe}] }
        """
        seance    = self.get_object()
        presences = request.data.get('presences', [])
        errors    = []
        count     = 0
        # Seuls les élèves inscrits dans la classe de la séance peuvent être pointés.
        inscrits = set(Inscription.objects.filter(code_classe=seance.code_classe_id, code_annee=seance.code_annee_id)
                       .values_list('mle_etudiant_id', flat=True))
        for item in presences:
            mle = item.get('mle_etudiant')
            if not mle:
                errors.append({'mle_etudiant': None, 'detail': 'Matricule manquant.'})
                continue
            if mle not in inscrits:
                errors.append({'mle_etudiant': mle, 'detail': "Élève non inscrit dans cette classe."})
                continue
            try:
                Absence.objects.update_or_create(
                    code_seance=seance,
                    mle_etudiant_id=mle,
                    defaults={
                        'present':   item.get('present', False),
                        'signe':     item.get('signe', False),
                        'retard_minutes': max(int(item.get('retard_minutes') or 0), 0) if item.get('present') else 0,
                        # La justification (posée par le surveillant général) n'est modifiée
                        # que si elle est explicitement transmise : réenregistrer l'appel ne
                        # l'efface plus.
                        **{k: item[k] for k in ('motif', 'justifiee') if k in item},
                    }
                )
                count += 1
            except IntegrityError as exc:
                errors.append({'mle_etudiant': mle, 'detail': str(exc)})
        log_action(request, 'UPDATE', 'Seance', pk, f'{count} présences saisies')
        response_data = {'updated': count}
        if errors:
            response_data['errors'] = errors
        return Response(response_data)

    @action(detail=False, methods=['get'], url_path='export-csv')
    def export_csv(self, request):
        qs = self.filter_queryset(self.get_queryset())
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="seances.csv"'
        response.write('﻿')
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['Date', 'Matière', 'Classe', 'Enseignant', 'H. début', 'H. fin', 'Salle', 'Nb heures', 'Statut', 'Année'])
        for s in qs:
            writer.writerow([
                s.date_seance.strftime('%d/%m/%Y') if s.date_seance else '',
                str(s.code_matiere),
                str(s.code_classe),
                str(s.mle_ens) if s.mle_ens else '',
                s.h_debut.strftime('%H:%M') if s.h_debut else '',
                s.h_fin.strftime('%H:%M') if s.h_fin else '',
                s.salle or '',
                s.nb_heures_effectuees,
                s.statut,
                str(s.code_annee),
            ])
        return response

class AbsenceViewSet(EtablissementParRelationMixin, SMSBaseViewSet):
    etudiant_lookup = 'mle_etudiant'
    etablissement_lookups = ('mle_etudiant__etablissement',)
    queryset         = Absence.objects.select_related('code_seance', 'mle_etudiant').all()
    serializer_class = AbsenceSerializer
    search_fields    = ['mle_etudiant__nom']
    filterset_fields = ['code_seance', 'mle_etudiant', 'present', 'justifiee']


# ── Gestion interne ───────────────────────────────────────────────────────────
class ExamenViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    queryset         = Examen.objects.select_related(
        'code_matiere', 'code_classe', 'code_annee', 'code_periode').all()
    serializer_class = ExamenSerializer
    search_fields    = ['lib_examen', 'code_matiere__lib_matiere', 'code_classe__lib_classe', 'type_examen',
                        'type_officiel', 'organisme', 'surveillant', 'code_salle']
    filterset_fields = ['code_classe', 'code_annee', 'type_examen']

    def get_queryset(self):
        qs = super().get_queryset()
        # Un compte ENSEIGNANT ne planifie/voit que les examens de ses propres matières
        # et classes (dérivé de Cours) — pas ceux de toute l'école.
        return scope_queryset_to_enseignant_cours(self.request, qs)

    def perform_create(self, serializer):
        assert_enseignant_teaches(
            self.request,
            serializer.validated_data['code_matiere'].pk,
            serializer.validated_data['code_classe'].pk,
        )
        super().perform_create(serializer)

    def perform_update(self, serializer):
        code_matiere = serializer.validated_data.get('code_matiere', serializer.instance.code_matiere)
        code_classe  = serializer.validated_data.get('code_classe',  serializer.instance.code_classe)
        assert_enseignant_teaches(self.request, code_matiere.pk, code_classe.pk)
        super().perform_update(serializer)

    @action(detail=True, methods=['post'], url_path='generer-convocations')
    def generer_convocations(self, request, pk=None):
        examen = self.get_object()
        annee  = examen.code_annee
        inscriptions = Inscription.objects.filter(
            code_classe=examen.code_classe, code_annee=annee
        ).select_related('mle_etudiant')
        count = 0
        for insc in inscriptions:
            Convocation.objects.get_or_create(
                type_convocation='EXAMEN',
                destinataire_type='ETUDIANT',
                mle_etudiant=insc.mle_etudiant,
                code_examen=examen,
                defaults={
                    'objet':          f"Convocation — {examen.lib_examen}",
                    'date_evenement': examen.date_examen,
                    'lieu':           examen.code_salle or '',
                }
            )
            count += 1
        examen.convocation_envoyee = True
        examen.save()
        log_action(request, 'CREATE', 'Convocation', pk,
                   f'{count} convocations générées pour {examen.lib_examen}')
        return Response({'generated': count})

class ConvocationViewSet(EtablissementParRelationMixin, SMSBaseViewSet):
    etudiant_lookup = 'mle_etudiant'
    etablissement_lookups = ('mle_etudiant__etablissement', 'code_examen__etablissement')
    queryset         = Convocation.objects.select_related(
        'mle_etudiant', 'mle_ens', 'code_examen').all()
    serializer_class = ConvocationSerializer
    search_fields    = ['objet', 'mle_etudiant__nom']
    filterset_fields = ['type_convocation', 'destinataire_type', 'statut', 'mle_etudiant']


class EpreuveViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    """
    Sujets d'examen soumis par les enseignants pour validation par le service scolarité
    avant la date de l'examen (voir Epreuve.STATUT_CHOICES : BROUILLON → SOUMISE →
    VALIDEE/REJETEE). Un ENSEIGNANT ne voit/gère que ses propres épreuves ; SCOLARITE
    (et SUPER_ADMIN) voient toutes les épreuves de l'établissement et sont seuls habilités
    à valider/rejeter.
    """
    queryset         = Epreuve.objects.select_related(
        'examen__code_matiere', 'examen__code_classe', 'soumis_par', 'valide_par').all()
    serializer_class = EpreuveSerializer
    parser_classes   = [MultiPartParser, FormParser, JSONParser]
    filterset_fields = ['statut', 'examen', 'soumis_par']

    def get_queryset(self):
        qs = super().get_queryset()
        return scope_queryset_to_enseignant_cours(
            self.request, qs, matiere_field='examen__code_matiere', classe_field='examen__code_classe',
        )

    def _assert_can_validate(self):
        profile = get_utilisateur_profile(self.request)
        if not profile or profile.role not in ('SCOLARITE', 'SUPER_ADMIN'):
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied("Seul le service scolarité peut valider ou rejeter une épreuve.")

    def perform_create(self, serializer):
        examen = serializer.validated_data['examen']
        assert_enseignant_teaches(self.request, examen.code_matiere_id, examen.code_classe_id)
        extra = {}
        etab_id = self.get_etablissement_id()
        if etab_id:
            extra['etablissement_id'] = etab_id
        enseignant = get_enseignant_profile(self.request)
        if enseignant:
            extra['soumis_par'] = enseignant
        instance = serializer.save(**extra)
        log_action(self.request, 'CREATE', 'Epreuve', instance.pk)

    def perform_update(self, serializer):
        instance = serializer.instance
        if instance.statut in ('SOUMISE', 'VALIDEE') and not self._request_is_validator():
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied(
                "Cette épreuve est en attente de validation ou déjà validée : "
                "elle ne peut plus être modifiée directement."
            )
        super().perform_update(serializer)

    def _request_is_validator(self):
        profile = get_utilisateur_profile(self.request)
        return bool(profile and profile.role in ('SCOLARITE', 'SUPER_ADMIN'))

    @action(detail=True, methods=['post'], url_path='soumettre')
    def soumettre(self, request, pk=None):
        from django.utils import timezone
        epreuve = self.get_object()
        if epreuve.statut not in ('BROUILLON', 'REJETEE'):
            return Response(
                {'detail': "Cette épreuve a déjà été soumise ou validée."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if not epreuve.fichier:
            return Response(
                {'detail': "Aucun fichier n'a été déposé pour cette épreuve."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        epreuve.statut = 'SOUMISE'
        epreuve.date_soumission = timezone.now()
        epreuve.commentaire_validation = ''
        enseignant = get_enseignant_profile(request)
        if enseignant:
            epreuve.soumis_par = enseignant
        epreuve.save()
        log_action(request, 'UPDATE', 'Epreuve', pk, 'Soumission pour validation')
        return Response(self.get_serializer(epreuve).data)

    @action(detail=True, methods=['post'], url_path='valider')
    def valider(self, request, pk=None):
        from django.utils import timezone
        self._assert_can_validate()
        epreuve = self.get_object()
        if epreuve.statut != 'SOUMISE':
            return Response(
                {'detail': "Seule une épreuve soumise peut être validée."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        epreuve.statut = 'VALIDEE'
        epreuve.valide_par = get_utilisateur_profile(request)
        epreuve.date_validation = timezone.now()
        epreuve.commentaire_validation = ''
        epreuve.save()
        log_action(request, 'UPDATE', 'Epreuve', pk, 'Épreuve validée')
        return Response(self.get_serializer(epreuve).data)

    @action(detail=True, methods=['post'], url_path='rejeter')
    def rejeter(self, request, pk=None):
        from django.utils import timezone
        self._assert_can_validate()
        commentaire = (request.data.get('commentaire') or '').strip()
        if not commentaire:
            return Response(
                {'detail': "Un commentaire expliquant le rejet est requis."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        epreuve = self.get_object()
        if epreuve.statut != 'SOUMISE':
            return Response(
                {'detail': "Seule une épreuve soumise peut être rejetée."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        epreuve.statut = 'REJETEE'
        epreuve.valide_par = get_utilisateur_profile(request)
        epreuve.date_validation = timezone.now()
        epreuve.commentaire_validation = commentaire
        epreuve.save()
        log_action(request, 'UPDATE', 'Epreuve', pk, f'Épreuve rejetée : {commentaire}')
        return Response(self.get_serializer(epreuve).data)

def _scope_filter_kwargs(code_classe, code_sp, code_dep, code_faculte, classe_field='code_classe'):
    """
    Périmètre d'agrégation d'un rapport statistique, du plus précis au plus large :
    classe > spécialité > filière (département) > pôle (faculté). Un seul niveau est
    retenu — c'est celui-là qui délimite le groupe agrégé (ex : choisir une spécialité
    agrège toutes les classes de cette spécialité, tous départements confondus le cas
    échéant, pas seulement celles du département actuellement sélectionné).
    `classe_field` est le chemin vers la classe depuis le modèle filtré (ex : 'code_classe'
    pour Inscription/Decision/Evaluation, qui portent chacun une FK classe directe).
    """
    if code_classe:
        return {classe_field: code_classe}
    if code_sp:
        return {f'{classe_field}__code_sp': code_sp}
    if code_dep:
        return {f'{classe_field}__code_dep': code_dep}
    if code_faculte:
        return {f'{classe_field}__code_dep__code_faculte': code_faculte}
    return {}


class RapportStatistiqueViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    queryset         = RapportStatistique.objects.select_related(
        'code_annee', 'code_classe', 'code_dep', 'code_sp', 'code_faculte').all()
    serializer_class = RapportStatistiqueSerializer
    filterset_fields = ['code_annee', 'code_classe', 'code_dep', 'code_sp', 'code_faculte']

    @action(detail=False, methods=['post'], url_path='generate')
    def generate(self, request):
        code_annee   = request.data.get('code_annee')
        code_classe  = request.data.get('code_classe')  or None
        code_sp      = request.data.get('code_sp')      or None
        code_dep     = request.data.get('code_dep')     or None
        code_faculte = request.data.get('code_faculte') or None

        if not code_annee:
            return Response({'error': 'code_annee est requis'}, status=status.HTTP_400_BAD_REQUEST)

        scope = _scope_filter_kwargs(code_classe, code_sp, code_dep, code_faculte)

        # ── Inscriptions ────────────────────────────────────────────────────────
        inscr_qs = Inscription.objects.filter(code_annee=code_annee).filter(**scope)

        nb_inscrits = inscr_qs.count()
        nb_hommes   = inscr_qs.filter(mle_etudiant__sexe='M').count()
        nb_femmes   = inscr_qs.filter(mle_etudiant__sexe='F').count()

        # ── Décisions ───────────────────────────────────────────────────────────
        dec_qs = Decision.objects.filter(code_annee=code_annee).filter(**scope)

        nb_admis     = dec_qs.filter(resultat='ADMIS').count()
        nb_ajournes  = dec_qs.filter(resultat='AJOURNE').count()
        nb_redoubles = dec_qs.filter(resultat='REDOUBLE').count()

        # ── Évaluations ─────────────────────────────────────────────────────────
        eval_qs = Evaluation.objects.filter(**scope)

        agg = eval_qs.aggregate(avg=Avg('note'), mn=Min('note'), mx=Max('note'))
        moyenne  = round(float(agg['avg']), 2) if agg['avg'] is not None else None
        note_min = round(float(agg['mn']),  2) if agg['mn']  is not None else None
        note_max = round(float(agg['mx']),  2) if agg['mx']  is not None else None

        genere_par = request.user.get_full_name() or request.user.username \
            if request.user.is_authenticated else 'Système'

        rapport = RapportStatistique.objects.create(
            code_annee_id      = code_annee,
            code_classe_id     = code_classe,
            code_sp_id         = code_sp,
            code_dep_id        = code_dep,
            code_faculte_id    = code_faculte,
            nb_inscrits        = nb_inscrits,
            nb_hommes          = nb_hommes,
            nb_femmes          = nb_femmes,
            nb_admis           = nb_admis,
            nb_ajournes        = nb_ajournes,
            nb_redoubles       = nb_redoubles,
            moyenne_generale   = moyenne,
            note_min           = note_min,
            note_max           = note_max,
            genere_par         = genere_par,
            # Sans ceci, le rapport reste invisible pour tout compte non SUPER_ADMIN
            # (EtablissementFilterMixin.get_queryset filtre par etablissement_id) — cette
            # action crée l'objet directement plutôt que via perform_create, donc doit
            # reproduire son affectation d'établissement elle-même.
            etablissement_id   = self.get_etablissement_id(),
        )
        serializer = self.get_serializer(rapport)
        return Response(serializer.data, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['get'], url_path='export-csv')
    def export_csv(self, request):
        qs = self.filter_queryset(self.get_queryset())
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="rapport_statistique.csv"'
        response.write('﻿')
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['Année', 'Classe', 'Spécialité', 'Département', 'Pôle', 'Inscrits', 'Hommes', 'Femmes',
                         'Admis', 'Ajournés', 'Taux réussite %', 'Taux féminisation %', 'Moyenne'])
        for r in qs:
            writer.writerow([
                r.code_annee.lib_annee if r.code_annee else '',
                r.code_classe.lib_classe if r.code_classe else '',
                r.code_sp.lib_sp if r.code_sp else '',
                r.code_dep.lib_dep if r.code_dep else '',
                r.code_faculte.lib_faculte if r.code_faculte else '',
                r.nb_inscrits, r.nb_hommes, r.nb_femmes,
                r.nb_admis, r.nb_ajournes,
                r.taux_reussite, r.taux_feminisation,
                r.moyenne_generale or '',
            ])
        return response


class RapportAssiduiteViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    interdit_etudiant = True  # données réservées au personnel
    queryset         = RapportAssiduite.objects.select_related(
        'code_annee', 'code_periode', 'code_classe', 'code_sp', 'code_dep', 'code_faculte').all()
    serializer_class = RapportAssiduiteSerializer
    filterset_fields = ['code_annee', 'code_periode', 'code_classe', 'code_sp', 'code_dep', 'code_faculte']

    @action(detail=False, methods=['post'], url_path='generate')
    def generate(self, request):
        code_periode_id = request.data.get('code_periode')
        code_classe     = request.data.get('code_classe')  or None
        code_sp         = request.data.get('code_sp')      or None
        code_dep        = request.data.get('code_dep')     or None
        code_faculte    = request.data.get('code_faculte') or None

        if not code_periode_id:
            return Response({'error': 'code_periode est requis'}, status=status.HTTP_400_BAD_REQUEST)
        periode = Periode.objects.select_related('code_annee').filter(pk=code_periode_id).first()
        if not periode:
            return Response({'error': 'Période introuvable'}, status=status.HTTP_400_BAD_REQUEST)
        if not periode.code_annee_id:
            return Response(
                {'error': "Cette période n'est rattachée à aucune année scolaire."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        scope = _scope_filter_kwargs(
            code_classe, code_sp, code_dep, code_faculte, classe_field='code_seance__code_classe',
        )

        # ── Présences/absences enregistrées dans le périmètre et la période ──────
        abs_qs = Absence.objects.filter(
            code_seance__code_annee_id=periode.code_annee_id,
        ).filter(**scope)
        if periode.date_debut:
            abs_qs = abs_qs.filter(code_seance__date_seance__gte=periode.date_debut.date())
        if periode.date_fin:
            abs_qs = abs_qs.filter(code_seance__date_seance__lte=periode.date_fin.date())

        nb_controles             = abs_qs.count()
        nb_absences              = abs_qs.filter(present=False).count()
        nb_absences_injustifiees = abs_qs.filter(present=False, justifiee=False).count()
        nb_seances               = abs_qs.values('code_seance').distinct().count()
        nb_etudiants             = abs_qs.values('mle_etudiant').distinct().count()

        genere_par = request.user.get_full_name() or request.user.username \
            if request.user.is_authenticated else 'Système'

        rapport = RapportAssiduite.objects.create(
            code_annee_id             = periode.code_annee_id,
            code_periode_id           = periode.pk,
            code_classe_id            = code_classe,
            code_sp_id                = code_sp,
            code_dep_id               = code_dep,
            code_faculte_id           = code_faculte,
            periode_debut             = periode.date_debut.date() if periode.date_debut else None,
            periode_fin               = periode.date_fin.date()   if periode.date_fin   else None,
            nb_etudiants              = nb_etudiants,
            nb_seances                = nb_seances,
            nb_controles              = nb_controles,
            nb_absences               = nb_absences,
            nb_absences_injustifiees  = nb_absences_injustifiees,
            genere_par                = genere_par,
            etablissement_id          = self.get_etablissement_id(),
        )
        serializer = self.get_serializer(rapport)
        return Response(serializer.data, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['get'], url_path='export-csv')
    def export_csv(self, request):
        qs = self.filter_queryset(self.get_queryset())
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="rapport_assiduite.csv"'
        response.write('﻿')
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['Année', 'Période', 'Classe', 'Spécialité', 'Département', 'Pôle',
                         'Étudiants', 'Séances', 'Contrôles', 'Absences', 'Absences injustifiées',
                         'Taux absence %', 'Taux présence %'])
        for r in qs:
            writer.writerow([
                r.code_annee.lib_annee if r.code_annee else '',
                r.code_periode.lib_periode if r.code_periode else '',
                r.code_classe.lib_classe if r.code_classe else '',
                r.code_sp.lib_sp if r.code_sp else '',
                r.code_dep.lib_dep if r.code_dep else '',
                r.code_faculte.lib_faculte if r.code_faculte else '',
                r.nb_etudiants, r.nb_seances, r.nb_controles,
                r.nb_absences, r.nb_absences_injustifiees,
                r.taux_absence, r.taux_presence,
            ])
        return response


# ── Documents avancés ─────────────────────────────────────────────────────────
class DecisionViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    etudiant_lookup = 'mle_etudiant'
    queryset         = Decision.objects.select_related(
        'mle_etudiant', 'code_annee', 'code_classe').all()
    serializer_class = DecisionSerializer
    search_fields    = ['mle_etudiant__nom']
    filterset_fields = ['code_annee', 'code_classe', 'resultat', 'mention', 'session']

    def get_queryset(self):
        qs = super().get_queryset()
        code_dep = self.request.query_params.get('code_dep')
        code_sp  = self.request.query_params.get('code_sp')
        if code_dep:
            qs = qs.filter(code_classe__code_dep=code_dep)
        if code_sp:
            qs = qs.filter(code_classe__code_sp=code_sp)
        return qs

    ROLES_DELIBERATION = {'SUPER_ADMIN', 'ADMIN', 'SCOLARITE', 'CENSEUR'}

    @action(detail=False, methods=['post'], url_path='generer')
    def generer(self, request):
        """
        Délibération automatique d'une classe pour une année (voir notes.deliberer) :
        crée ou met à jour la décision de chaque élève noté. Body : code_classe,
        code_annee, session (NORMALE|RATTRAPAGE, défaut NORMALE), seuil (défaut 10).
        """
        from rest_framework.exceptions import PermissionDenied
        profile = get_utilisateur_profile(request)
        if not profile or profile.role not in self.ROLES_DELIBERATION:
            raise PermissionDenied("Votre rôle ne vous autorise pas à générer les décisions.")

        code_classe = request.data.get('code_classe')
        code_annee  = request.data.get('code_annee')
        session     = request.data.get('session') or 'NORMALE'
        if not code_classe or not code_annee:
            return Response({'error': 'code_classe et code_annee sont requis.'},
                            status=status.HTTP_400_BAD_REQUEST)
        if session not in dict(Decision.SESSION_CHOICES):
            return Response({'error': 'Session invalide.'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            seuil = float(request.data.get('seuil', 10))
        except (TypeError, ValueError):
            return Response({'error': 'Seuil invalide.'}, status=status.HTTP_400_BAD_REQUEST)
        # La classe doit être visible pour l'utilisateur (isolation multi-établissement).
        if not _scoped_queryset(ClasseViewSet, request).filter(pk=code_classe).exists():
            return Response({'error': 'Classe introuvable.'}, status=status.HTTP_404_NOT_FOUND)

        evaluations = _scoped_queryset(EvaluationViewSet, request)
        bilan = deliberer(evaluations, code_classe, code_annee, session=session, seuil=seuil)
        log_action(request, 'CREATE', 'Decision', code_classe,
                   f"Délibération {session} {code_annee} : {bilan['crees']} créées, "
                   f"{bilan['mis_a_jour']} mises à jour")
        return Response(bilan)

    @action(detail=False, methods=['get'], url_path='export-csv')
    def export_csv(self, request):
        qs = self.filter_queryset(self.get_queryset())
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="decisions.csv"'
        response.write('﻿')
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['Matricule', 'Nom', 'Classe', 'Année', 'Session',
                         'Moyenne', 'Crédits validés', 'Résultat', 'Mention', 'Rang'])
        for d in qs:
            writer.writerow([
                d.mle_etudiant.mle_etudiant, d.mle_etudiant.nom,
                d.code_classe.lib_classe if d.code_classe else '',
                d.code_annee.lib_annee  if d.code_annee  else '',
                d.get_session_display(),
                d.moyenne_annuelle or '', d.credits_valides,
                d.get_resultat_display(), d.mention or '', d.rang or '',
            ])
        return response

class DiplomeViewSet(EtablissementParRelationMixin, SMSBaseViewSet):
    etudiant_lookup = 'mle_etudiant'
    etablissement_lookups = ('mle_etudiant__etablissement',)
    queryset         = Diplome.objects.select_related('mle_etudiant').all()
    serializer_class = DiplomeSerializer
    search_fields    = ['lib_diplome', 'mle_etudiant__nom', 'numero_serie']
    filterset_fields = ['annee_obtention', 'mention', 'type_diplome']

class CarteEtudiantViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    etudiant_lookup = 'mle_etudiant'
    queryset         = CarteEtudiant.objects.select_related('mle_etudiant').all()
    serializer_class = CarteEtudiantSerializer
    search_fields    = ['numero_carte', 'mle_etudiant__nom', 'mle_etudiant__prenom', 'mle_etudiant__mle_etudiant', 'statut', 'code_annee__code_annee', 'code_annee__lib_annee']
    filterset_fields = ['code_annee', 'statut', 'mle_etudiant']

class StageViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    etudiant_lookup = 'mle_etudiant'
    queryset         = Stage.objects.select_related('mle_etudiant').all()
    serializer_class = StageSerializer
    search_fields    = ['entreprise', 'sujet', 'tuteur_entreprise', 'type_stage', 'statut', 'mle_etudiant__nom', 'mle_etudiant__prenom', 'mle_etudiant__mle_etudiant']
    filterset_fields = ['code_annee', 'type_stage', 'statut', 'certificat_emis']

    def get_queryset(self):
        qs = super().get_queryset()
        code_dep = self.request.query_params.get('code_dep')
        code_sp  = self.request.query_params.get('code_sp')
        if code_dep:
            qs = qs.filter(mle_etudiant__code_dep=code_dep)
        if code_sp:
            qs = qs.filter(mle_etudiant__code_sp=code_sp)
        return qs

    @action(detail=False, methods=['get'], url_path='export-csv')
    def export_csv(self, request):
        qs = self.filter_queryset(self.get_queryset())
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="stages.csv"'
        response.write('﻿')
        writer = csv.writer(response, delimiter=';')
        writer.writerow(['Matricule', 'Étudiant', 'Entreprise', 'Sujet', 'Type', 'Début', 'Fin', 'Note', 'Statut', 'Certificat', 'Année'])
        for s in qs:
            e = s.mle_etudiant
            writer.writerow([
                e.mle_etudiant,
                f'{e.nom} {e.prenom or ""}'.strip(),
                s.entreprise,
                s.sujet or '',
                s.get_type_stage_display(),
                s.date_debut.strftime('%d/%m/%Y') if s.date_debut else '',
                s.date_fin.strftime('%d/%m/%Y') if s.date_fin else '',
                s.note_stage if s.note_stage is not None else '',
                s.get_statut_display(),
                'Oui' if s.certificat_emis else 'Non',
                str(s.code_annee),
            ])
        return response

class LettreAdmissionViewSet(EtablissementParRelationMixin, SMSBaseViewSet):
    etudiant_lookup = 'mle_etudiant'
    etablissement_lookups = ('mle_etudiant__etablissement',)
    queryset         = LettreAdmission.objects.select_related('mle_etudiant').all()
    serializer_class = LettreAdmissionSerializer
    search_fields    = ['mle_etudiant__nom']
    filterset_fields = ['type_lettre', 'code_annee', 'envoye_par_email']

class BadgeAccesViewSet(EtablissementParRelationMixin, SMSBaseViewSet):
    etudiant_lookup = 'mle_etudiant'
    etablissement_lookups = ('mle_etudiant__etablissement', 'mle_ens__etablissement')
    queryset         = BadgeAcces.objects.select_related('mle_etudiant', 'mle_ens').all()
    serializer_class = BadgeAccesSerializer
    search_fields    = ['numero_badge', 'mle_etudiant__nom']
    filterset_fields = ['type_porteur', 'statut']

class DocumentGenereViewSet(EtablissementParRelationMixin, SMSBaseViewSet):
    etudiant_lookup = 'mle_etudiant'
    etablissement_lookups = ('mle_etudiant__etablissement', 'code_classe__etablissement')
    queryset         = DocumentGenere.objects.select_related(
        'mle_etudiant', 'code_classe', 'code_annee').all()
    serializer_class = DocumentGenereSerializer
    search_fields    = ['nom_fichier', 'type_document']
    filterset_fields = ['type_document', 'format', 'mle_etudiant', 'code_annee']


# ── Multi-établissement ───────────────────────────────────────────────────────
class NiveauScolaireViewSet(SMSBaseViewSet):
    queryset         = NiveauScolaire.objects.all()
    serializer_class = NiveauScolaireSerializer
    filterset_fields = ['type_etab', 'systeme']
    search_fields    = ['lib_niveau', 'lib_en']

class CoefficientMatiereViewSet(EtablissementFilterMixin, SMSBaseViewSet):
    """
    Grille des coefficients (primaire / secondaire) de l'établissement courant : un
    coefficient par année, niveau, série et matière — voir CoefficientMatiere et
    notes.coefficients_classe. Écriture réservée à la direction pédagogique.
    """
    queryset         = CoefficientMatiere.objects.select_related('code_matiere').all()
    serializer_class = CoefficientMatiereSerializer
    filterset_fields = ['code_annee', 'niveau_scolaire', 'serie', 'code_matiere']
    ROLES_ECRITURE   = {'SUPER_ADMIN', 'ADMIN', 'SCOLARITE', 'CENSEUR'}

    def check_permissions(self, request):
        super().check_permissions(request)
        if request.method not in ('GET', 'HEAD', 'OPTIONS'):
            from rest_framework.exceptions import PermissionDenied
            profile = get_utilisateur_profile(request)
            if not profile or profile.role not in self.ROLES_ECRITURE:
                raise PermissionDenied("Votre rôle ne vous autorise pas à modifier les coefficients.")

    def _etablissement_courant(self):
        """École courante (celle du compte, ou celle choisie par le SUPER_ADMIN)."""
        from rest_framework.exceptions import ValidationError
        etab_id = self.get_etablissement_id()
        if not etab_id:
            raise ValidationError({'error': 'Aucun établissement sélectionné.'})
        if type_etab_courant(self.request) not in ('PRIMAIRE', 'SECONDAIRE'):
            raise ValidationError({'error': "Les coefficients ne concernent que le primaire et le secondaire."})
        return etab_id

    @action(detail=False, methods=['get'], url_path='grille')
    def grille(self, request):
        """
        Grille d'un niveau pour une année : séries en colonnes, matières en lignes.
        Matières proposées : celles enseignées (Cours) dans les classes de ce niveau,
        plus celles déjà présentes dans la grille. Query : code_annee, niveau_scolaire.
        """
        code_annee = request.query_params.get('code_annee')
        niveau     = request.query_params.get('niveau_scolaire')
        if not code_annee or not niveau:
            return Response({'error': 'code_annee et niveau_scolaire sont requis.'},
                            status=status.HTTP_400_BAD_REQUEST)
        etab_id = self._etablissement_courant()
        classes = Classe.objects.filter(etablissement_id=etab_id, niveau_scolaire_id=niveau)
        enseignees = set(Cours.objects.filter(code_classe__in=classes, code_annee=code_annee)
                         .values_list('code_matiere', flat=True))
        lignes = CoefficientMatiere.objects.filter(
            etablissement_id=etab_id, code_annee=code_annee, niveau_scolaire_id=niveau)
        coefs = {}
        for l in lignes:
            coefs.setdefault(l.code_matiere_id, {})[l.serie] = l.coef
        matieres = Matiere.objects.filter(code_matiere__in=enseignees | set(coefs)).order_by('lib_matiere')
        return Response({
            'series':   series_du_niveau(niveau),
            'classes':  list(classes.order_by('lib_classe').values('code_classe', 'lib_classe', 'serie')),
            'matieres': [{'code_matiere': m.code_matiere, 'lib_matiere': m.lib_matiere,
                          'enseignee': m.code_matiere in enseignees,
                          'coefs': coefs.get(m.code_matiere, {})} for m in matieres],
        })

    @action(detail=False, methods=['post'], url_path='enregistrer-grille')
    def enregistrer_grille(self, request):
        """
        Enregistre la grille d'un niveau en une fois. Body : code_annee, niveau_scolaire,
        lignes [{code_matiere, serie, coef}] — un coef vide supprime la case.
        """
        from django.db import transaction
        code_annee = request.data.get('code_annee')
        niveau     = request.data.get('niveau_scolaire')
        lignes     = request.data.get('lignes') or []
        if not code_annee or not niveau:
            return Response({'error': 'code_annee et niveau_scolaire sont requis.'},
                            status=status.HTTP_400_BAD_REQUEST)
        etab_id = self._etablissement_courant()
        series_valides = set(series_du_niveau(niveau))
        erreurs, a_ecrire = [], []
        for l in lignes:
            serie = l.get('serie') or ''
            coef  = l.get('coef')
            if serie not in series_valides:
                erreurs.append(f"Série {serie!r} invalide pour ce niveau.")
                continue
            if coef in (None, ''):
                a_ecrire.append((l.get('code_matiere'), serie, None))
                continue
            try:
                coef = int(coef)
            except (TypeError, ValueError):
                coef = 0
            if not 1 <= coef <= 20:
                erreurs.append(f"Coefficient invalide pour {l.get('code_matiere')} {serie} : entre 1 et 20.")
                continue
            a_ecrire.append((l.get('code_matiere'), serie, coef))
        if erreurs:
            return Response({'error': ' '.join(erreurs)}, status=status.HTTP_400_BAD_REQUEST)

        crees = modifies = supprimes = 0
        with transaction.atomic():
            for code_matiere, serie, coef in a_ecrire:
                cle = dict(etablissement_id=etab_id, code_annee_id=code_annee,
                           niveau_scolaire_id=niveau, serie=serie, code_matiere_id=code_matiere)
                if coef is None:
                    supprimes += CoefficientMatiere.objects.filter(**cle).delete()[0]
                    continue
                _, cree = CoefficientMatiere.objects.update_or_create(**cle, defaults={'coef': coef})
                crees += cree
                modifies += not cree
        log_action(request, 'UPDATE', 'CoefficientMatiere', niveau,
                   f'Grille {niveau} {code_annee} : {crees} créés, {modifies} modifiés, {supprimes} supprimés')
        return Response({'crees': crees, 'modifies': modifies, 'supprimes': supprimes})

    @action(detail=False, methods=['post'], url_path='copier')
    def copier(self, request):
        """
        Copie la grille d'une année vers une autre (ex : année précédente → nouvelle
        année). Les cases déjà remplies dans l'année cible ne sont pas écrasées.
        Body : source_annee, cible_annee.
        """
        source, cible = request.data.get('source_annee'), request.data.get('cible_annee')
        if not source or not cible or source == cible:
            return Response({'error': 'source_annee et cible_annee (différentes) sont requises.'},
                            status=status.HTTP_400_BAD_REQUEST)
        etab_id = self._etablissement_courant()
        existantes = set(CoefficientMatiere.objects.filter(etablissement_id=etab_id, code_annee=cible)
                         .values_list('niveau_scolaire', 'serie', 'code_matiere'))
        nouvelles = [
            CoefficientMatiere(etablissement_id=etab_id, code_annee_id=cible, niveau_scolaire_id=c.niveau_scolaire_id,
                               serie=c.serie, code_matiere_id=c.code_matiere_id, coef=c.coef)
            for c in CoefficientMatiere.objects.filter(etablissement_id=etab_id, code_annee=source)
            if (c.niveau_scolaire_id, c.serie, c.code_matiere_id) not in existantes
        ]
        CoefficientMatiere.objects.bulk_create(nouvelles)
        log_action(request, 'CREATE', 'CoefficientMatiere', cible, f'{len(nouvelles)} coefficients copiés depuis {source}')
        return Response({'copies': len(nouvelles)})

class ConfigBulletinViewSet(SMSBaseViewSet):
    queryset         = ConfigBulletin.objects.all()
    serializer_class = ConfigBulletinSerializer
    filterset_fields = ['type_etab']


# ── Envoi d'email ─────────────────────────────────────────────────────────────
class EnvoiEmailView(APIView):
    """
    POST /api/envoi-email/
    Body: { "to": "...", "subject": "...", "body_html": "...", "body_text": "..." }
    Envoie un email via le backend SMTP configuré dans settings.py.
    """
    # Sans ceci, n'importe quel compte authentifié (même ETUDIANT) pouvait faire relayer
    # un email arbitraire par le SMTP de l'établissement.
    permission_classes = [IsAuthenticated, RoleBasedPermission]

    def post(self, request):
        to        = (request.data.get('to') or '').strip()
        subject   = (request.data.get('subject') or 'SMS — Notification').strip()
        body_html = request.data.get('body_html', '')
        body_text = request.data.get('body_text') or body_html

        if not to or '@' not in to:
            return Response({'error': 'Adresse email invalide ou manquante.'},
                            status=status.HTTP_400_BAD_REQUEST)

        from_email = getattr(django_settings, 'DEFAULT_FROM_EMAIL', 'noreply@sms-ecole.cm')
        try:
            send_mail(
                subject=subject,
                message=body_text,
                from_email=from_email,
                recipient_list=[to],
                html_message=body_html or None,
                fail_silently=False,
            )
            log_action(request, 'UPDATE', 'Email', None, f'Email envoyé à {to} — {subject}')
            return Response({'sent': True, 'to': to})
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


class DashboardStatsView(APIView):
    """
    GET /api/dashboard/stats/
    Renvoie tous les KPIs + séries temporelles du tableau de bord en une seule requête.
    Remplace les ~14 appels séparés du frontend par ~12 requêtes SQL optimisées (COUNT/SUM/GROUP BY).
    """
    permission_classes = [IsAuthenticated]

    def _etab_filter(self, request):
        """Réplique la logique de EtablissementFilterMixin pour les APIView simples."""
        etab_id = trusted_etablissement_header(request)
        if not etab_id:
            # École indéterminée : refus par défaut, jamais de vue « toutes écoles ».
            return {'etablissement_id__in': []}
        return {'etablissement_id': etab_id} if etab_id else {}

    def get(self, request):
        import logging, traceback
        logger = logging.getLogger(__name__)
        try:
            return self._compute(request)
        except Exception as exc:
            logger.error('DashboardStatsView error: %s\n%s', exc, traceback.format_exc())
            return Response({'error': str(exc), 'type': type(exc).__name__}, status=500)

    def _compute(self, request):
        from django.db.models import Count, Sum, Q
        from django.db.models.functions import ExtractYear, ExtractMonth
        from datetime import date

        f = self._etab_filter(request)

        # ── Comptages simples ─────────────────────────────────────────────────
        nb_etudiants    = Etudiant.objects.filter(**f).count()
        nb_enseignants  = Enseignant.objects.filter(**f).count()
        nb_classes      = Classe.objects.filter(**f).count()
        nb_inscriptions = Inscription.objects.filter(**f).count()
        nb_seances      = Seance.objects.filter(**f).count()
        nb_decisions    = Decision.objects.filter(**f).count()
        nb_stages       = Stage.objects.filter(statut='EN_COURS', **f).count()
        nb_examens      = Examen.objects.filter(**f).count()

        # ── Évaluations : taux de réussite + distribution notes (1 query) ────
        eval_agg = Evaluation.objects.filter(**f).aggregate(
            total=Count('code_eval'),
            reussies=Count('code_eval', filter=Q(note__gte=10)),
            lt8=Count('code_eval',    filter=Q(note__lt=8)),
            n8_10=Count('code_eval',  filter=Q(note__gte=8,  note__lt=10)),
            n10_12=Count('code_eval', filter=Q(note__gte=10, note__lt=12)),
            n12_14=Count('code_eval', filter=Q(note__gte=12, note__lt=14)),
            n14_16=Count('code_eval', filter=Q(note__gte=14, note__lt=16)),
            n16_20=Count('code_eval', filter=Q(note__gte=16)),
        )
        taux_reussite = (
            round(eval_agg['reussies'] / eval_agg['total'] * 100)
            if eval_agg['total'] else None
        )

        # ── Paiements : nombre d'impayés (1 query) ────────────────────────────
        paiement_agg = Paiement.objects.filter(**f).aggregate(
            nb_impayes=Count('code_paiement', filter=Q(statut='IMPAYE')),
        )

        # ── Factures : solde + taux de recouvrement (1 query) ─────────────────
        facture_agg = Facture.objects.filter(**f).aggregate(
            nb_factures=Count('code_facture'),
            nb_soldees=Count('code_facture', filter=Q(statut='SOLDEE')),
            total=Sum('montant_total'),
            total_paye=Sum('montant_paye'),
        )
        nb_factures  = facture_agg['nb_factures'] or 0
        nb_soldees   = facture_agg['nb_soldees']  or 0
        solde_impaye = float((facture_agg['total'] or 0) - (facture_agg['total_paye'] or 0))
        taux_rec     = round(nb_soldees / nb_factures * 100) if nb_factures else None

        # ── Répartition par département + comptage étudiants (1 query) ────────
        deps = list(
            Departement.objects.filter(**f)
            .annotate(nb_et=Count('etudiant'))
            .values('code_dep', 'lib_dep', 'nb_et')
            .order_by('-nb_et')[:8]
        )

        # ── Séries temporelles — 8 derniers mois (2 queries) ──────────────────
        today  = date.today()
        cutoff = date(today.year, today.month, 1)
        # Reculer de 7 mois pour obtenir le premier des 8 mois
        m, y = cutoff.month - 7, cutoff.year
        if m <= 0:
            m += 12
            y -= 1
        cutoff = date(y, m, 1)

        paiements_mois = list(
            Paiement.objects.filter(date_paiement__gte=cutoff, **f)
            .annotate(annee=ExtractYear('date_paiement'), mois_num=ExtractMonth('date_paiement'))
            .values('annee', 'mois_num')
            .annotate(total=Sum('mt_paiement'))
            .order_by('annee', 'mois_num')
        )
        inscriptions_mois = list(
            Inscription.objects.filter(date_inscription__gte=cutoff, **f)
            .annotate(annee=ExtractYear('date_inscription'), mois_num=ExtractMonth('date_inscription'))
            .values('annee', 'mois_num')
            .annotate(count=Count('code_inscription'))
            .order_by('annee', 'mois_num')
        )

        return Response({
            # ── Comptages ──
            'nb_etudiants':    nb_etudiants,
            'nb_enseignants':  nb_enseignants,
            'nb_classes':      nb_classes,
            'nb_inscriptions': nb_inscriptions,
            'nb_evaluations':  eval_agg['total'] or 0,
            'nb_seances':      nb_seances,
            'nb_decisions':    nb_decisions,
            'nb_stages':       nb_stages,
            'nb_examens':      nb_examens,
            'nb_factures':     nb_factures,
            # ── KPIs calculés ──
            'nb_impayes':        paiement_agg['nb_impayes'] or 0,
            'solde_impaye':      solde_impaye,
            'taux_reussite':     taux_reussite,
            'taux_recouvrement': taux_rec,
            # ── Graphiques ──
            'repartition_deps': [
                {'code_dep': d['code_dep'], 'lib_dep': d['lib_dep'], 'nb_et': d['nb_et']}
                for d in deps
            ],
            'paiements_mois': [
                {'mois': f"{p['annee']}-{str(p['mois_num']).zfill(2)}", 'total': float(p['total'] or 0)}
                for p in paiements_mois
            ],
            'inscriptions_mois': [
                {'mois': f"{p['annee']}-{str(p['mois_num']).zfill(2)}", 'count': p['count']}
                for p in inscriptions_mois
            ],
            'dist_notes': [
                eval_agg.get('lt8',   0) or 0,
                eval_agg.get('n8_10', 0) or 0,
                eval_agg.get('n10_12',0) or 0,
                eval_agg.get('n12_14',0) or 0,
                eval_agg.get('n14_16',0) or 0,
                eval_agg.get('n16_20',0) or 0,
            ],
        })


# ── Import CSV ────────────────────────────────────────────────────────────────

def _resolve_etab(request):
    """
    Retourne l'objet Etablissement pour la requête courante.
    SUPER_ADMIN peut cibler un établissement précis via le corps (code_etab) ou le header
    X-Etablissement-Id. Tout autre rôle importe toujours dans son propre établissement,
    quoi qu'il envoie dans le corps ou les headers — sinon n'importe quel compte pourrait
    importer des données dans une autre école (IDOR en écriture).
    """
    from .models import Etablissement
    if is_super_admin(request):
        target = request.data.get('code_etab') or trusted_etablissement_header(request)
        if target:
            return Etablissement.objects.filter(code_etab=target).first()
        return Etablissement.objects.filter(actif=True).first()
    etab_id = trusted_etablissement_header(request)
    return Etablissement.objects.filter(code_etab=etab_id).first() if etab_id else None


class ImportCsvView(APIView):
    """
    POST /api/import-csv/<entity_type>/
    body (multipart/form-data):
      - file       : fichier CSV
      - dry_run    : "true" | "false"  (défaut "false")
      - code_etab  : code d'établissement cible (SUPER_ADMIN uniquement)
    """
    # Import en masse — même restriction d'écriture que le reste de l'API (voir
    # RoleBasedPermission) : ETUDIANT/DIRECTION/APEE ne peuvent pas déclencher d'import.
    permission_classes = [IsAuthenticated, RoleBasedPermission]
    parser_classes     = [MultiPartParser, FormParser]

    def post(self, request, entity_type):
        from .imports import IMPORTERS

        if entity_type not in IMPORTERS:
            return Response(
                {'error': f"Type inconnu : '{entity_type}'. "
                          f"Valeurs acceptées : {', '.join(IMPORTERS.keys())}"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        csv_file = request.FILES.get('file')
        if not csv_file:
            return Response({'error': 'Aucun fichier envoyé.'}, status=status.HTTP_400_BAD_REQUEST)

        dry_run_raw = request.data.get('dry_run', 'false')
        dry_run     = str(dry_run_raw).lower() in ('1', 'true', 'yes', 'oui')

        etab = _resolve_etab(request)
        if not etab:
            return Response(
                {'error': "Établissement introuvable ou non déterminable pour cet utilisateur."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        importer = IMPORTERS[entity_type]()
        result   = importer.run(csv_file, etab, dry_run=dry_run)

        if not dry_run and result['success'] > 0 and not result['errors']:
            detail = (
                f"{entity_type}: {result['created']} créé(s), "
                f"{result['updated']} mis à jour — "
                f"{result['total']} ligne(s) traitée(s)"
            )
            log_action(
                request, 'IMPORT_CSV', entity_type,
                etab.code_etab if etab else '',
                detail,
            )

        return Response(result, status=status.HTTP_200_OK)


class ImportCsvTemplateView(APIView):
    """
    GET /api/import-csv/<entity_type>/template/
    Télécharge le gabarit Excel (.xlsx) mis en forme pour l'entité.
    """
    permission_classes = [IsAuthenticated]

    def get(self, request, entity_type):
        from .imports import IMPORTERS
        import io as _io

        if entity_type not in IMPORTERS:
            return Response(
                {'error': f"Type inconnu : '{entity_type}'."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        importer = IMPORTERS[entity_type]()
        xlsx_buf = importer.generate_template_xlsx()
        response = HttpResponse(
            xlsx_buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = (
            f'attachment; filename="gabarit_{entity_type}.xlsx"'
        )
        return response

